import io
import json
import math
import re
from urllib.parse import quote

import pandas as pd
import plotly.graph_objects as go
import requests
import streamlit as st

# ---------------------------------------------------------------
# FREE ENGINES
#   1) Europe PMC REST API  - free, no key (literature + hit counts)
#   2) Google Gemini API    - free tier key from aistudio.google.com (optional)
#      No key? The app still runs with rule-based estimates.
# ---------------------------------------------------------------
EPMC_URL = "https://www.ebi.ac.uk/europepmc/webservices/rest/search"
GEMINI_URL = "https://generativelanguage.googleapis.com/v1beta/models/{model}:generateContent"
MODEL_FALLBACKS = ["gemini-flash-latest", "gemini-3.5-flash", "gemini-3-flash-preview", "gemini-3.1-flash-lite"]
PILLARS = ["Novelty", "Feasibility", "Pathway", "Methodology", "Work Plan"]

STOP = set("""
a an the of in on at for to and or with by from among between within into using use used
study studies evaluating evaluation assess assessing assessment analysis analyzing effect effects
impact role based cross sectional cross-sectional retrospective prospective observational this that their than then also over under about
""".split())


# ---------------------- file reading ----------------------
def read_upload(up):
    """Return plain text from .txt/.csv/.pdf/.docx uploads."""
    if up is None:
        return ""
    name, data = up.name.lower(), up.getvalue()
    try:
        if name.endswith(".pdf"):
            from pypdf import PdfReader
            reader = PdfReader(io.BytesIO(data))
            return "\n".join((pg.extract_text() or "") for pg in reader.pages)
        if name.endswith(".docx"):
            from docx import Document
            doc = Document(io.BytesIO(data))
            parts = [p.text for p in doc.paragraphs]
            for table in doc.tables:
                for row in table.rows:
                    parts.append(" | ".join(c.text for c in row.cells))
            return "\n".join(parts)
    except ImportError:
        st.error("Missing library. In Terminal run: python3 -m pip install pypdf python-docx")
        return ""
    except Exception as e:
        st.error(f"Could not read '{up.name}': {e}")
        return ""
    return data.decode("utf-8", errors="ignore")


# ---------------------- evidence retrieval ----------------------
@st.cache_data(ttl=3600, show_spinner=False)
def epmc_search(query, n=10):
    r = requests.get(
        EPMC_URL,
        params={"query": query, "format": "json", "resultType": "core", "pageSize": n},
        timeout=20,
    )
    r.raise_for_status()
    data = r.json()
    papers = []
    for p in data.get("resultList", {}).get("result", []):
        papers.append({
            "Title": (p.get("title") or "").rstrip("."),
            "Year": p.get("pubYear"),
            "PMID": p.get("pmid"),
            "DOI": p.get("doi"),
            "Abstract": (p.get("abstractText") or "")[:900],
        })
    return int(data.get("hitCount", 0)), papers


def extract_terms(text, limit=6):
    words = re.findall(r"[A-Za-z][A-Za-z\-]{3,}", text.lower())
    seen, terms = set(), []
    for w in words:
        if w not in STOP and w not in seen:
            seen.add(w)
            terms.append(w)
    return terms[:limit]


def _clean(t):
    return re.sub(r'["\\]', "", str(t)).strip().lower()[:60]


def simple_words(text, limit=5):
    words = re.findall(r"[A-Za-z0-9][A-Za-z0-9\-]+", text.lower())
    out = []
    for w in words:
        if w not in STOP and w not in out:
            out.append(w)
    return out[:limit]


def parse_keywords(text, is_list=False):
    """Turn a few typed keywords (or a sentence) into concept dicts. Commas keep phrases together."""
    text = text.strip()
    if is_list or len(text.split()) <= 10:
        parts = [p.strip() for p in re.split(r"[,;\n]+", text) if p.strip()]
        if len(parts) < 2:
            parts = simple_words(text)
    else:
        parts = extract_terms(text, limit=5)
    concepts = []
    for p in parts:
        name = _clean(p)
        if name and name not in [c["name"] for c in concepts]:
            concepts.append({"name": name, "synonyms": []})
    return concepts[:5]


def concept_clause(c):
    terms = [c["name"]] + c["synonyms"]
    inner = " OR ".join(f'TITLE_ABS:"{t}"' for t in terms)
    return f"({inner})" if len(terms) > 1 else inner


def build_query(concepts):
    """Title/abstract only (Europe PMC searches full text by default, which inflates hit counts)."""
    return " AND ".join(concept_clause(c) for c in concepts)


def title_overlap(concepts, papers):
    """Best % of your concepts that appear together in one retrieved title."""
    best = 0.0
    for p in papers:
        t = p["Title"].lower()
        hit = sum(1 for c in concepts if any(w in t for w in [c["name"]] + c["synonyms"]))
        best = max(best, 100.0 * hit / len(concepts))
    return round(best)


def hit_novelty(hits):
    """0 hits -> 100% ; ~10,000 hits -> 0%."""
    return max(0.0, 100.0 - 100.0 * math.log10(hits + 1) / 4.0)


# ---------------------- rule-based fallback ----------------------
def rule_scores(text, has_draft):
    t = text.lower()
    human = any(w in t for w in ["patient", "participant", "survey", "questionnaire",
                                 "practitioner", "doctor", "clinician", "nurse", "resident", "student", "blood", "serum", "interview", "cohort"])
    lit_only = any(w in t for w in ["scoping", "systematic review", "meta-analysis", "in silico",
                                    "bioinformatic", "database", "literature"])
    wet_lab = any(w in t for w in ["assay", "culture", "pcr", "gene", "knockout", "cell line", "animal"])
    design = any(w in t for w in ["cross-sectional", "cohort", "case-control", "scoping", "systematic",
                                  "rct", "randomi", "survey", "in vitro", "in silico", "prisma"])
    sample = any(w in t for w in ["sample", "n=", "population", "participants", "tertiary", "hospital", "centre", "center"])
    outcome = any(w in t for w in ["prevalence", "incidence", "adherence", "awareness", "score", "outcome", "level", "compliance"])
    analysis = any(w in t for w in ["regression", "chi-square", "spss", "thematic", "prisma", "anova", "t-test"])
    timeline = any(w in t for w in ["week", "month", "timeline", "phase", "milestone"])
    ethics_word = any(w in t for w in ["consent", "ethic", "iec", "irb"])

    feas = 85 - (25 if human else 0) - (20 if wet_lab else 0) + (10 if lit_only else 0)
    path = (90 if lit_only and not human else 55 if human else 70) + (10 if ethics_word else 0)
    meth = 15 + 25 * design + 20 * sample + 20 * outcome + 20 * analysis
    plan = 35 + 35 * timeline + (15 if has_draft else 0) + (15 if design else 0)
    clamp = lambda v: int(max(5, min(100, v)))
    notes = {
        "Feasibility": "Human participants detected - needs recruitment + IEC." if human else "No participant recruitment detected.",
        "Pathway": "Ethics/regulatory path inferred from keywords only.",
        "Methodology": "Scored on whether design, population, outcome and analysis are mentioned.",
        "Work Plan": "Scored on whether a timeline/milestones are mentioned.",
    }
    return {"Feasibility": clamp(feas), "Pathway": clamp(path),
            "Methodology": clamp(meth), "Work Plan": clamp(plan)}, notes, human


# ---------------------- free LLM (Gemini) ----------------------
def _gemini_json(key, model, prompt):
    body = {"contents": [{"parts": [{"text": prompt}]}],
            "generationConfig": {"responseMimeType": "application/json", "temperature": 0.2}}
    tried = []
    for m in dict.fromkeys([model] + MODEL_FALLBACKS):
        r = requests.post(GEMINI_URL.format(model=m),
                          headers={"x-goog-api-key": key, "Content-Type": "application/json"},
                          json=body, timeout=90)
        if r.status_code == 404:  # model retired / renamed -> try the next one
            tried.append(m)
            continue
        if not r.ok:
            raise RuntimeError(f"Gemini error {r.status_code}: {r.text[:200]}")
        txt = r.json()["candidates"][0]["content"]["parts"][0]["text"]
        return json.loads(txt.replace("```json", "").replace("```", "").strip()), m
    raise RuntimeError(f"No Gemini model found (tried: {', '.join(tried)}). Set a current model name in the sidebar.")


def plan_search(key, model, text, custom):
    """AI step: work out what a few rough keywords most likely mean, and expand them into search concepts."""
    prompt = f"""You are a biomedical literature-search expert helping an Indian MBBS student.
They typed a few keywords or a rough idea. Work out what they most likely mean.
INPUT: {text[:1500]}
USER KEYWORDS: {custom}

Return ONLY JSON:
{{"interpreted_proposal": "one or two clear sentences describing the most likely research proposal",
 "study_design_guess": "",
 "concepts": [{{"name": "", "synonyms": ["", ""]}}],
 "note": "what is ambiguous or assumed"}}
Rules: 2-5 concepts ordered from most central to most specific; 2-5 synonyms each (alternate spellings,
abbreviations, MeSH-style terms, e.g. 'heart attack' -> 'myocardial infarction'); keep each term 1-4 words;
do not add concepts the user did not imply."""
    data, used = _gemini_json(key, model, prompt)
    concepts = []
    for c in data.get("concepts", [])[:5]:
        if not isinstance(c, dict):
            continue
        name = _clean(c.get("name", ""))
        syns = [_clean(x) for x in (c.get("synonyms") or [])[:5]]
        syns = [x for x in dict.fromkeys(syns) if x and x != name]
        if name:
            concepts.append({"name": name, "synonyms": syns})
    if not concepts:
        raise ValueError("planner returned no concepts")
    return {"interpreted": str(data.get("interpreted_proposal", ""))[:600],
            "design": str(data.get("study_design_guess", ""))[:200],
            "note": str(data.get("note", ""))[:400], "concepts": concepts}, used


def gemini_analyze(key, model, proposal, papers, stats, extra_text):
    ev = [{"pmid": p["PMID"], "year": p["Year"], "title": p["Title"], "abstract": p["Abstract"]} for p in papers]
    prompt = f"""You review medical research proposals for Indian MBBS students (ICMR STS style).
Use ONLY the evidence below. Cite PMIDs from the evidence list only. Never invent papers or numbers.
If evidence is thin, say so in "limits".
If EXTRA is a draft, review it in "draft_review" and cross-check its claims against the evidence; otherwise leave draft_review empty.

PROPOSAL: {proposal}
SEARCH STATS: {json.dumps(stats)}
EVIDENCE: {json.dumps(ev)}
EXTRA (guideline/draft text): {extra_text[:12000]}

Return ONLY JSON:
{{"scores": {{"Novelty": 0-100, "Feasibility": 0-100, "Pathway": 0-100, "Methodology": 0-100, "Work Plan": 0-100}},
 "notes": {{"Novelty": "", "Feasibility": "", "Pathway": "", "Methodology": "", "Work Plan": ""}},
 "already_done": ["short items with PMID"],
 "gap": "what is still unexplored",
 "methodology": {{"design": "", "population": "", "sample_size_approach": "", "tools": "", "analysis": ""}},
 "work_plan": [{{"weeks": "", "task": ""}}],
 "ethics": {{"human_participants": true, "likely_category": "exempt / expedited / full board", "reasons": ""}},
 "draft_review": {{"strengths": [], "problems": [], "contradictions": []}},
 "limits": ""}}"""
    return _gemini_json(key, model, prompt)


# ---------------------- chart ----------------------
def radar(scores):
    vals = [scores[p] for p in PILLARS]
    labels = [f"<b>{p}</b><br>{scores[p]}%" for p in PILLARS]
    fig = go.Figure(go.Scatterpolar(
        r=vals + vals[:1], theta=labels + labels[:1],
        fill="toself", mode="lines+markers",
    ))
    fig.update_layout(
        polar=dict(
            gridshape="linear",
            radialaxis=dict(range=[0, 100], tickvals=[20, 40, 60, 80, 100], ticksuffix="%"),
            angularaxis=dict(rotation=90, direction="clockwise"),
        ),
        showlegend=False, height=470, margin=dict(l=70, r=70, t=30, b=30),
    )
    return fig


def verdict(n):
    if n >= 75:
        return "Highly novel - confirm no unindexed duplicates"
    if n >= 50:
        return "Moderately novel - sharpen your unique angle"
    if n >= 25:
        return "Crowded topic - narrow the population, setting or outcome"
    return "Heavily studied - consider a different angle"


def build_report(proposal, scores, notes, queries, stats, ai, audit=None):
    lines = ["RESEARCH IDEA ANALYSIS", "", f"Proposal: {proposal}", "", "SCORES"]
    lines += [f"- {p}: {scores[p]}%  ({notes.get(p, '')})" for p in PILLARS]
    lines += ["", "SEARCHES"] + [f"- {q} -> {s['hits']} hits" for q, s in zip(queries, stats)]
    if ai:
        lines += ["", f"GAP: {ai.get('gap', '')}", f"METHODOLOGY: {json.dumps(ai.get('methodology', {}))}",
                  f"ETHICS: {json.dumps(ai.get('ethics', {}))}", f"LIMITS: {ai.get('limits', '')}"]
        lines += [f"- {w.get('weeks')}: {w.get('task')}" for w in ai.get("work_plan", [])]
    if audit:
        lines += ["", f"DRAFT CHECK: {audit['words']} words (limit {audit['limit'] or 'not set'})",
                  "Missing sections: " + (", ".join(k for k, v in audit["sections"].items() if not v) or "none"),
                  "Placeholders: " + (", ".join(audit["placeholders"]) or "none")]
        lines += [f"- {d['DOI']}: {d['Status']}" for d in audit["dois"]]
    lines += ["", "Advisory only. Final ethics decision rests with your IEC."]
    return "\n".join(lines)


# ---------------------- uploaded draft check ----------------------
SECTIONS = {
    "Aim / objectives": ["aim", "objective"],
    "Methodology": ["methodology", "methods", "study design"],
    "Sample size / population": ["sample size", "population", "inclusion"],
    "Ethics / consent": ["ethic", "consent", "iec", "irb"],
    "Timeline": ["timeline", "week", "month", "milestone"],
    "References": ["reference", "doi", "et al"],
}
PLACEHOLDERS = ["tbd", "to be decided", "pending", "unconfirmed", "xx", "lorem", "[insert", "todo"]


@st.cache_data(ttl=86400, show_spinner=False)
def check_doi(doi):
    try:
        r = requests.get("https://api.crossref.org/works/" + quote(doi, safe="/"),
                         headers={"User-Agent": "ResearchIdeaAnalyzer/1.0"}, timeout=10)
        if r.status_code == 200:
            return "Found", ((r.json().get("message", {}).get("title") or [""])[0])
        if r.status_code == 404:
            return "Not in Crossref (typo, or a non-Crossref DOI)", ""
        return f"Check failed ({r.status_code})", ""
    except Exception:
        return "Check failed (network)", ""


def audit_draft(text, word_limit):
    t = text.lower()
    words = len(text.split())
    sections = {k: any(w in t for w in v) for k, v in SECTIONS.items()}
    holders = [h for h in PLACEHOLDERS if re.search(r"(?<!\w)" + re.escape(h) + r"(?!\w)", t)]
    dois = sorted({d.rstrip(".,;)]") for d in re.findall(r"10\.\d{4,9}/[^\s\"<>]+", text)})[:15]
    doi_rows = [{"DOI": d, "Status": (c := check_doi(d))[0], "Registered title": c[1]} for d in dois]
    return {"words": words, "limit": word_limit, "sections": sections, "placeholders": holders, "dois": doi_rows}


# ---------------------- analysis ----------------------
def run_analysis(proposal, custom, key, model, extra, word_limit=0):
    plan = {"interpreted": "", "design": "", "note": "", "error": None, "ai": False}
    concepts = None
    if key:
        try:
            with st.spinner("Interpreting your idea..."):
                planned, _ = plan_search(key, model, proposal, custom)
            concepts = planned.pop("concepts")
            plan.update(planned, ai=True)
        except Exception as e:
            plan["error"] = str(e)
    if not concepts:
        concepts = parse_keywords(custom, is_list=True) if custom.strip() else parse_keywords(proposal)
    if not concepts:
        raise ValueError("Could not find any keywords - type a few words describing your idea.")
    if plan["ai"] and plan["interpreted"] and len(proposal.split()) < 25:
        proposal = plan["interpreted"]  # downstream scoring works on the clearer wording

    n = len(concepts)
    ks = sorted({min(2, n), min(4, n), n})
    queries = [build_query(concepts[:k]) for k in ks]
    labels = [" AND ".join(c["name"] for c in concepts[:k]) for k in ks]
    mid = len(ks) // 2

    stats, papers = [], []
    with st.status("Searching Europe PMC...", expanded=True) as status:
        for i, q in enumerate(queries):
            hits, found = epmc_search(q, 10 if i == mid else 1)
            stats.append({"query": labels[i], "hits": hits, "novelty": round(hit_novelty(hits))})
            if i == mid:
                papers = found
            status.write(f"Search {i + 1}: `{labels[i]}` -> **{hits:,}** papers")
        if not papers:
            _, papers = epmc_search(queries[0], 10)
        status.update(label="Search complete", state="complete", expanded=False)

    sim = title_overlap(concepts, papers)
    main_stat = stats[mid]
    novelty = int(round(0.5 * main_stat["novelty"] + 0.5 * (100 - sim)))

    rule, rule_notes, _ = rule_scores(proposal + " " + extra, bool(extra))
    scores = {"Novelty": novelty, **rule}
    notes = {"Novelty": f"{main_stat['hits']:,} title/abstract matches for '{main_stat['query']}'; "
                        f"best single title covers {sim}% of your keywords.", **rule_notes}
    source, ai, ai_error = "Rule-based estimates (add a Gemini key for AI assessment)", None, None

    if key:
        try:
            with st.spinner("AI assessing the evidence..."):
                ai, used = gemini_analyze(key, model, proposal, papers, stats, extra)
            new_scores = {p: int(max(0, min(100, ai["scores"][p]))) for p in PILLARS[1:]}
            new_notes = {p: ai.get("notes", {}).get(p, "") for p in PILLARS[1:]}
            scores.update(new_scores)
            notes.update(new_notes)
            scores["Novelty"] = int(round(0.7 * novelty + 0.3 * ai["scores"]["Novelty"]))
            notes["Novelty"] += " " + ai.get("notes", {}).get("Novelty", "")
            source = f"AI assessment ({used}) grounded on retrieved papers"
        except Exception as e:
            ai, ai_error = None, str(e)

    audit = None
    if extra.strip():
        with st.spinner("Checking your uploaded draft..."):
            audit = audit_draft(extra, word_limit)

    return {"proposal": proposal, "scores": scores, "audit": audit, "plan": plan,
            "concepts": concepts, "queries": queries, "mid": mid, "notes": notes, "source": source, "ai": ai,
            "ai_error": ai_error, "stats": stats, "papers": papers, "labels": labels,
            "report": build_report(proposal, scores, notes, labels, stats, ai, audit)}


def render(res, framework):
    scores, notes, ai = res["scores"], res["notes"], res["ai"]
    plan = res.get("plan") or {}
    with st.expander("🧠 How I understood your idea", expanded=True):
        if plan.get("error"):
            st.warning(f"AI interpretation failed, used your words directly. ({plan['error']})")
        elif not plan.get("ai"):
            st.caption("No Gemini key: keywords were used as typed, without synonym expansion.")
        if plan.get("interpreted"):
            st.markdown(f"**Interpreted proposal:** {plan['interpreted']}")
        if plan.get("design"):
            st.markdown(f"**Likely design:** {plan['design']}")
        if plan.get("note"):
            st.caption(plan["note"])
        st.dataframe(pd.DataFrame([{"Concept": c["name"], "Also searched as": ", ".join(c["synonyms"]) or "-"}
                                   for c in res["concepts"]]), hide_index=True)
        st.caption("Europe PMC query behind the novelty score:")
        st.code(res["queries"][res["mid"]], language=None)

    st.subheader("📊 Pentagon Assessment")
    st.caption(res["source"])
    if res["ai_error"]:
        st.warning(f"AI step failed, showing rule-based scores. ({res['ai_error']})")
    st.plotly_chart(radar(scores))

    st.subheader("🔎 Novelty progress")
    for s in res["stats"]:
        st.progress(s["novelty"] / 100, text=f"{s['query']} -> {s['hits']:,} papers -> {s['novelty']}% novel")
    st.metric("Overall novelty", f"{scores['Novelty']}%")
    st.info(verdict(scores["Novelty"]))

    st.subheader("📋 Pillar details")
    st.dataframe(pd.DataFrame({"Pillar": PILLARS, "Score %": [scores[p] for p in PILLARS],
                               "Basis": [notes.get(p, "") for p in PILLARS]}), hide_index=True)

    if ai:
        st.markdown(f"**Gap:** {ai.get('gap', '')}")
        if ai.get("already_done"):
            st.markdown("**Already done:**\n" + "\n".join(f"- {x}" for x in ai["already_done"]))
        st.markdown("**Methodology**")
        st.json(ai.get("methodology", {}))
        st.markdown("**Work plan**")
        st.dataframe(pd.DataFrame(ai.get("work_plan", [])), hide_index=True)
        st.markdown("**Ethics screen**")
        st.json(ai.get("ethics", {}))
        dr = ai.get("draft_review") or {}
        if any(dr.get(k) for k in ("strengths", "problems", "contradictions")):
            st.markdown("**AI review of your draft**")
            for k, icon in (("strengths", "✅"), ("problems", "⚠️"), ("contradictions", "🚨")):
                for item in dr.get(k, []):
                    st.markdown(f"{icon} {item}")
        if ai.get("limits"):
            st.caption(f"Limits: {ai['limits']}")

    audit = res.get("audit")
    if audit:
        st.subheader("📁 Uploaded draft check")
        if audit["limit"]:
            over = audit["words"] > audit["limit"]
            st.metric("Word count", f"{audit['words']:,} / {audit['limit']:,}",
                      delta="over limit" if over else "within limit", delta_color="inverse")
        else:
            st.metric("Word count", f"{audit['words']:,}")
        st.dataframe(pd.DataFrame({"Section": list(audit["sections"]),
                                   "Found in draft": ["✅" if v else "⚠️ missing" for v in audit["sections"].values()]}),
                     hide_index=True)
        if audit["placeholders"]:
            st.warning("Placeholder text found: " + ", ".join(audit["placeholders"]))
        if audit["dois"]:
            st.markdown("**DOIs cited in your draft (checked against Crossref)**")
            st.dataframe(pd.DataFrame(audit["dois"]), hide_index=True)
        else:
            st.caption("No DOIs found in the draft to verify.")

    st.subheader("📚 Closest papers found")
    dfp = pd.DataFrame(res["papers"])
    if not dfp.empty:
        dfp["Link"] = dfp["DOI"].apply(lambda d: f"https://doi.org/{d}" if d else "")
        st.dataframe(dfp[["Title", "Year", "PMID", "Link"]], hide_index=True)
    else:
        st.write("No papers retrieved.")

    st.download_button("📥 Download report (.txt)", res["report"], file_name="research_analysis.txt")
    st.caption(f"Advisory only ({framework}). Novelty depends on the search terms - read the top papers yourself. "
               "Ethics/exemption decisions rest with your IEC.")


# ---------------------- app ----------------------
def main():
    st.set_page_config(page_title="Research Idea Analyzer", page_icon="🧠", layout="centered")
    st.title("🧠 Research Idea Analyzer")
    st.caption("Describe your proposed research. The app searches the literature (Europe PMC, free), "
               "measures novelty, and plots a 5-pillar pentagon.")

    with st.sidebar:
        st.header("⚙️ Free engine")
        key = st.text_input("Gemini API key (free at aistudio.google.com)", type="password")
        if not key:
            try:
                key = st.secrets.get("GEMINI_API_KEY", "")
            except Exception:
                key = ""
        model = st.text_input("Gemini model", value="gemini-flash-latest")
        st.caption("No key = rule-based estimates for 4 pillars. Novelty is always computed from real search data.")
        framework = st.selectbox("Target", ["ICMR STS", "International journal", "General academic"])
        up = st.file_uploader("Optional: guideline or draft", type=["txt", "csv", "pdf", "docx"])
        word_limit = st.number_input("Word limit to check the draft against (0 = off)", min_value=0, value=300, step=50)
        st.caption("Set this from your call document - I haven't verified the official figure.")
        extra = read_upload(up)
        if up and extra.strip():
            st.caption(f"Read {len(extra.split()):,} words from {up.name}")
        elif up:
            st.warning("No text found (scanned PDFs are images and can't be read).")

    proposal = st.text_area("1. What are you proposing? (or leave blank and upload your draft)", height=140,
                            placeholder="A few keywords is enough (e.g. beta cell, insulin, implant) "
                                        "or a full sentence")
    custom = st.text_input("2. (Optional) Your own keywords, comma-separated (use commas to keep phrases together)")

    if st.button("🚀 Analyze", type="primary"):
        prop = proposal.strip() or extra[:1500].strip()
        if len(prop) < 3:
            st.warning("Type a few keywords or a sentence about your idea, or upload a draft with text.")
        else:
            try:
                st.session_state["result"] = run_analysis(prop, custom, key, model, extra, word_limit)
            except Exception as e:
                st.session_state.pop("result", None)
                st.error(f"Analysis failed: {e}")

    if "result" in st.session_state:  # kept in session so the download button doesn't wipe the page
        render(st.session_state["result"], framework)


main()
