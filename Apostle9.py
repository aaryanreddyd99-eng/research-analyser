import io
import json
import math
import re
from pathlib import Path
from xml.sax.saxutils import escape
from urllib.parse import quote

import pandas as pd
import plotly.graph_objects as go
import requests
import streamlit as st

# ---------------------------------------------------------------
# FREE ENGINES
#   1) Europe PMC REST API  - free, no key (literature + hit counts)
#   2) Google Gemini API    - free tier key from aistudio.google.com (optional)
#      No key? The app still runs with rule-based estimates.
# ---------------------------------------------------------------
EPMC_URL = "https://www.ebi.ac.uk/europepmc/webservices/rest/search"
GEMINI_URL = "https://generativelanguage.googleapis.com/v1beta/models/{model}:generateContent"
MODEL_FALLBACKS = ["gemini-flash-latest", "gemini-3.5-flash", "gemini-3-flash-preview", "gemini-3.1-flash-lite"]
PILLARS = ["Novelty", "Feasibility", "Pathway", "Methodology", "Work Plan"]

STOP = set("""
a an the of in on at for to and or with by from among between within into using use used
study studies evaluating evaluation assess assessing assessment analysis analyzing effect effects
impact role based cross sectional cross-sectional retrospective prospective observational this
scoping review systematic meta-analysis prisma prisma-scr survey questionnaire randomised randomized controlled trial that their than then also over under about
""".split())


# ===============================================================
# RESEARCH KNOWLEDGE BASE  (edit freely - the app reads this)
# Curated general reference, NOT ethics/legal advice. Programme rules change
# every cycle: always verify against the current call document.
# ===============================================================
DESIGNS = {
    "Cross-sectional survey": {
        "keywords": ["cross-sectional", "cross sectional", "survey", "prevalence", "questionnaire", "kap study"],
        "guideline": "STROBE (cross-sectional checklist)",
        "sample_size": "Single proportion: n = Z^2 * p(1-p) / d^2 (Cochran); apply a finite-population correction for small sampling frames",
        "analysis": "Descriptive statistics; chi-square / t-test / logistic regression for associations",
        "ethics": "Involves participants, so IEC review is needed. Anonymous low-risk surveys are often expedited or exempt - the IEC decides",
        "fit": "High", "ethics_base": 60,
        "must_mention": {"Sample size": ["sample size"], "Inclusion/exclusion criteria": ["inclusion", "exclusion", "eligib"],
                         "Data collection tool": ["questionnaire", "instrument", "proforma", "tool"],
                         "Statistical analysis": ["spss", "statistic", "chi-square", "regression", "analysis"],
                         "Consent": ["consent"]},
    },
    "Case-control study": {
        "keywords": ["case-control", "case control"],
        "guideline": "STROBE (case-control checklist)",
        "sample_size": "From exposure prevalence in controls, expected odds ratio, power (usually 80%) and alpha 0.05",
        "analysis": "Odds ratio with 95% CI; logistic regression (conditional if matched)",
        "ethics": "Uses participants or patient records, so IEC review is needed; retrospective record review may be expedited - the IEC decides",
        "fit": "Medium", "ethics_base": 55,
        "must_mention": {"Case definition": ["case definition", "cases were", "definition of case"],
                         "Control selection": ["control"], "Sample size": ["sample size"],
                         "Confounders/matching": ["confound", "matched", "matching"], "Statistical analysis": ["odds ratio", "regression", "analysis"]},
    },
    "Cohort study": {
        "keywords": ["cohort", "follow-up", "follow up", "prospective", "retrospective"],
        "guideline": "STROBE (cohort checklist)",
        "sample_size": "From expected incidence in exposed vs unexposed, relative risk, power and attrition",
        "analysis": "Incidence, relative risk / hazard ratio, Kaplan-Meier, Cox regression",
        "ethics": "Participants are followed up, so IEC review is needed",
        "fit": "Low", "ethics_base": 50,
        "must_mention": {"Follow-up duration": ["follow-up", "follow up", "duration"], "Exposure definition": ["exposure", "exposed"],
                         "Outcome definition": ["outcome", "endpoint"], "Sample size": ["sample size"],
                         "Attrition": ["attrition", "loss to follow", "dropout"]},
    },
    "Randomized controlled trial": {
        "keywords": ["randomi", "rct", "randomised", "randomized", "clinical trial", "placebo"],
        "guideline": "CONSORT",
        "sample_size": "Two-group comparison formula using effect size, power, alpha and a dropout allowance",
        "analysis": "Intention-to-treat; t-test / chi-square / ANCOVA as appropriate",
        "ethics": "Full committee review; trial registration (CTRI in India) is expected before enrolment",
        "fit": "Low", "ethics_base": 30,
        "must_mention": {"Randomisation method": ["randomi"], "Allocation concealment/blinding": ["blind", "conceal", "mask"],
                         "Sample size": ["sample size"], "Primary outcome": ["primary outcome", "primary endpoint"],
                         "Trial registration": ["ctri", "registr"]},
    },
    "Systematic review / meta-analysis": {
        "keywords": ["systematic review", "meta-analysis", "meta analysis", "prospero"],
        "guideline": "PRISMA 2020 (PROSPERO registration is optional but encouraged)",
        "sample_size": "No sample size: define databases, search strings, PICO eligibility and number of reviewers",
        "analysis": "Risk-of-bias tool (e.g. RoB 2, Newcastle-Ottawa), narrative synthesis or random-effects meta-analysis, heterogeneity (I^2)",
        "ethics": "No participant contact, so usually exempt - the IEC decides",
        "fit": "Medium", "ethics_base": 90,
        "must_mention": {"Databases searched": ["pubmed", "database", "embase", "scopus"], "Search strategy": ["search strategy", "search string", "keywords"],
                         "Eligibility criteria": ["inclusion", "eligib", "pico"], "Risk of bias/quality appraisal": ["risk of bias", "quality assessment", "newcastle", "rob 2"],
                         "PRISMA": ["prisma"]},
    },
    "Scoping review": {
        "keywords": ["scoping review", "prisma-scr", "arksey", "evidence map"],
        "guideline": "PRISMA-ScR (Arksey & O'Malley / JBI framework)",
        "sample_size": "No sample size: define databases, search strings, eligibility concept/context/population",
        "analysis": "Data-charting table with narrative or thematic mapping; formal quality appraisal is optional",
        "ethics": "No participant contact, so usually exempt - the IEC decides",
        "fit": "High", "ethics_base": 90,
        "must_mention": {"Databases searched": ["pubmed", "database", "scopus", "embase"], "Search strategy": ["search strategy", "search string", "keywords"],
                         "Eligibility criteria": ["inclusion", "eligib", "concept"], "Data charting/extraction": ["charting", "extraction", "data extract"],
                         "PRISMA-ScR": ["prisma"]},
    },
    "Qualitative study": {
        "keywords": ["qualitative", "focus group", "in-depth interview", "thematic analysis", "phenomenolog"],
        "guideline": "COREQ or SRQR",
        "sample_size": "Purposive sampling until data saturation",
        "analysis": "Thematic or content analysis with coding; show trustworthiness (triangulation, member checking)",
        "ethics": "IEC review with informed consent and a confidentiality plan",
        "fit": "Medium", "ethics_base": 55,
        "must_mention": {"Sampling approach": ["purposive", "sampling", "saturation"], "Interview/focus-group guide": ["interview guide", "focus group", "interview"],
                         "Analysis method": ["thematic", "content analysis", "coding"], "Consent": ["consent"]},
    },
    "Laboratory (in vitro) study": {
        "keywords": ["in vitro", "cell line", "assay", "cell culture", "pcr", "western blot", "elisa"],
        "guideline": "Assay-specific reporting standards (e.g. MIQE for qPCR)",
        "sample_size": "Biological and technical replicates with controls (commonly at least triplicate)",
        "analysis": "t-test / ANOVA with post-hoc tests; dose-response curves",
        "ethics": "Human samples need IEC approval; institutional biosafety approval may apply",
        "fit": "Medium", "ethics_base": 70,
        "must_mention": {"Controls": ["control"], "Replicates": ["replicate", "triplicate", "repeated"], "Materials/reagents": ["reagent", "material", "kit"],
                         "Statistical analysis": ["anova", "t-test", "statistic"]},
    },
    "Animal study (in vivo)": {
        "keywords": ["animal", "rats", "mice", "mouse", "rodent", "in vivo", "wistar"],
        "guideline": "ARRIVE 2.0",
        "sample_size": "Power-based or resource-equation approach; apply the 3Rs (replace, reduce, refine)",
        "analysis": "t-test / ANOVA with randomisation and blinding reported",
        "ethics": "Institutional Animal Ethics Committee (IAEC / CPCSEA) approval is required",
        "fit": "Low", "ethics_base": 40,
        "must_mention": {"Species/strain": ["species", "strain", "wistar", "mice", "rats"], "IAEC approval": ["iaec", "cpcsea", "animal ethics"],
                         "Sample size": ["sample size", "n="], "Randomisation/blinding": ["randomi", "blind"]},
    },
    "In silico / bioinformatics": {
        "keywords": ["in silico", "bioinformatic", "docking", "computational", "database mining", "machine learning"],
        "guideline": "Report data sources, versions, tools and parameters so results are reproducible",
        "sample_size": "Dataset size and inclusion criteria from public databases",
        "analysis": "Statistical or computational validation, cross-validation, multiple-testing correction",
        "ethics": "Public anonymised data is usually exempt - the IEC decides",
        "fit": "High", "ethics_base": 90,
        "must_mention": {"Data sources": ["database", "dataset", "geo ", "tcga", "uniprot", "pdb"], "Tools/software": ["software", "tool", "python", "r package", "autodock"],
                         "Validation": ["validat", "cross-validation", "benchmark"]},
    },
    "Diagnostic accuracy study": {
        "keywords": ["diagnostic accuracy", "sensitivity", "specificity", "roc curve", "predictive value"],
        "guideline": "STARD",
        "sample_size": "From expected sensitivity/specificity, disease prevalence and desired precision",
        "analysis": "Sensitivity, specificity, PPV/NPV, ROC/AUC with 95% CI",
        "ethics": "Participants or records involved, so IEC review is needed",
        "fit": "Medium", "ethics_base": 55,
        "must_mention": {"Reference (gold) standard": ["gold standard", "reference standard"], "Index test": ["index test", "diagnostic test", "new test"],
                         "Sample size": ["sample size"], "Accuracy metrics": ["sensitivity", "specificity", "roc"]},
    },
    "Case report / case series": {
        "keywords": ["case report", "case series"],
        "guideline": "CARE",
        "sample_size": "Not applicable (report patient details with consent)",
        "analysis": "Descriptive",
        "ethics": "Written patient consent for publication; follow institutional policy on IEC notification",
        "fit": "High", "ethics_base": 65,
        "must_mention": {"Patient consent": ["consent"], "Timeline of events": ["timeline", "presented", "history"], "Literature review": ["literature", "review"]},
    },
}

ETHICS_CATEGORIES = {
    "Exemption from review": "Literature-only work or publicly available, anonymised data; no more than minimal risk",
    "Expedited review": "Minimal risk with participants or identifiable records (e.g. non-sensitive surveys, record review)",
    "Full committee review": "More than minimal risk, vulnerable groups, interventions or sensitive topics",
}
ETHICS_NOTE = ("General reading of the ICMR National Ethical Guidelines (2017). "
               "Only your Institutional Ethics Committee decides the category.")

# ICMR-DHR Short Term Studentship (STS) - taken from the official DHR documents read on 22 Oct 2026:
# schemes.dhr.gov.in/hrd -> 1.1 Concept Note format, 1.2 Topics and Subject Areas, 3.2 FAQ.
# Programme rules change between calls: always confirm on the portal (sts.icmr.org.in).
STS_RULES = [
    ("Only 1st and 2nd professional-year MBBS/BDS students of NMC/DCI-recognised Indian colleges can apply; 3rd/4th year, interns, PG, paramedical and foreign-college students cannot.", "DHR FAQ 3, 6, 11"),
    ("The guide must be a regular full-time faculty member of the same college; part-time, visiting faculty, residents, pool officers and PG students cannot guide.", "DHR FAQ 1, 10"),
    ("A guide can forward only one application per call; if more are forwarded, all are rejected.", "DHR FAQ 25"),
    ("A student can apply once per call during 1st/2nd year; once selected, they cannot take another STS in their lifetime.", "DHR FAQ 24"),
    ("Stipend Rs 60,000 in three instalments: Rs 15,000 after concept-note approval, Rs 15,000 after the data-collection report, Rs 30,000 after final-report approval.", "DHR FAQ 2"),
    ("The work runs over three consecutive summer breaks: (1) conceptualisation and research, (2) key research activities, (3) data analysis and final report.", "DHR FAQ 8, 17"),
    ("IEC approval is required if the project involves human participants or animal models. Submit the IEC certificate with the proposal or a guide undertaking that it will follow with the interim report.", "DHR FAQ 20"),
    ("Applications are fully online (no hard copies); the concept note must follow the DHR format.", "DHR FAQ 4, 13"),
    ("The project must be biomedical research; interdisciplinary approaches within biomedicine are acceptable.", "DHR FAQ 14"),
    ("Project and guide cannot be changed after selection except in exceptional, approved circumstances.", "DHR FAQ 16"),
    ("About 2,000 students are selected each year, roughly 80% MBBS and 20% BDS.", "DHR FAQ 7"),
]

# Concept-note format (DHR 1.1): (section, guide words, hard cap). The document says the guide may be
# exceeded by up to 100 words, except the title (25 words max).
STS_FORMAT = [("Title", 25, 25), ("Introduction", 300, 400), ("Objectives", 100, 200),
              ("Methodology", 800, 900), ("Implications / outcome", 100, 200), ("References", 300, 400)]
STS_FORMAT_NOTES = [
    "Times New Roman, size 12, single spacing, justified; STS Reference ID at the top left of page 1.",
    "No names, phone numbers or emails of student, guide or college anywhere - such proposals are rejected.",
    "For the data collection site write only the district, not the hospital or college name.",
    "References in Vancouver style, only from PubMed-indexed published articles.",
]
STS_TYPES = ["Clinical Investigations", "Epidemiological Investigations", "Field Operational Research", "Health Education",
             "Laboratory Investigations", "Product Development process", "Therapy and Management", "Any other (specify)"]
STS_SUBJECTS = ["Accident, Injuries & Trauma", "Allied Science", "Anatomy & Forensic medicine", "Anaesthesia",
                "Bioinformatics & Biostatistics", "Biochemistry", "Child Health/Pediatrics & its Subspecialties",
                "Critical & Intensive Care & Nursing", "Cardiology (includes surgical areas)", "Dermatology & Venereal Diseases",
                "Education & Advocacy", "Endocrinology", "Ear, Nose, Throat (ENT)", "Environment & Occupational Health",
                "Gastroenterology & Hepatology", "Geriatrics", "Gynecology & Obstetrics", "Hematology", "Human Genetics",
                "Microbiology/Immunology/Immunization", "Nephrology & Urology", "Neurology", "Obesity, Nutrition & Dietetics",
                "Oncology, Pathology & Cytology", "Ophthalmology", "Oral health & dentistry areas", "Orthopedics", "Pharmacology",
                "Physiology", "Public Health & Epidemiology", "Pulmonology & Respiratory Medicine",
                "Psychology, Mental Health & Socio-Behavioral Research", "Radiology", "Social & Preventive Medicine", "Surgery",
                "Traditional Medicine", "Any other (specify)"]

DEFAULT_JOURNALS = ["Indian Journal of Medical Research", "National Medical Journal of India", "Indian Journal of Public Health",
                    "BMJ Open", "PLoS One", "Cureus", "The Lancet Regional Health - Southeast Asia",
                    "Indian Journal of Pharmacology", "Journal of Family Medicine and Primary Care"]


def detect_design(text, ai_guess=""):
    t = (text + " " + ai_guess).lower()
    best, best_n = None, 0
    for name, d in DESIGNS.items():
        n = sum(t.count(k) for k in d["keywords"]) + (3 if name.lower().split(" /")[0] in ai_guess.lower() else 0)
        if n > best_n:
            best, best_n = name, n
    return best


def design_checks(name, text):
    if not name:
        return {}
    t = text.lower()
    return {label: any(w in t for w in words) for label, words in DESIGNS[name]["must_mention"].items()}


def build_reference(name, extra_rules="", framework="ICMR STS"):
    lines = ["CURATED REFERENCE (general; the user's own guideline text and pasted rules override it):"]
    if name:
        d = DESIGNS[name]
        lines += [f"Likely design: {name}", f"- Reporting guideline: {d['guideline']}", f"- Sample size approach: {d['sample_size']}",
                  f"- Typical analysis: {d['analysis']}", f"- Ethics: {d['ethics']}", f"- Fit for a short student project: {d['fit']}"]
    lines += ["Ethics categories (ICMR 2017, general): " + "; ".join(f"{k}: {v}" for k, v in ETHICS_CATEGORIES.items())]
    if framework == "ICMR STS":
        lines += ["ICMR STS rules (official DHR documents; confirm on the portal): " + " | ".join(f"{r} [{src}]" for r, src in STS_RULES)]
        lines += ["STS concept-note sections and word guides: " + ", ".join(f"{n} {g}" for n, g, _ in STS_FORMAT)]
        lines += ["STS types of study (choose exactly one): " + "; ".join(STS_TYPES)]
        lines += ["STS subject areas (choose exactly one): " + "; ".join(STS_SUBJECTS)]
    if extra_rules.strip():
        lines += ["RULES PASTED BY THE USER (highest priority): " + extra_rules.strip()[:3000]]
    return "\n".join(lines)


# ---------------------- file reading ----------------------
@st.cache_data(show_spinner=False)
def extract_text(name, data):
    """Plain text from .txt/.md/.csv/.pdf/.docx bytes. Raises on failure."""
    name = name.lower()
    if name.endswith(".pdf"):
        from pypdf import PdfReader
        reader = PdfReader(io.BytesIO(data))
        return "\n".join((pg.extract_text() or "") for pg in reader.pages)
    if name.endswith(".docx"):
        from docx import Document
        doc = Document(io.BytesIO(data))
        parts = [p.text for p in doc.paragraphs]
        for table in doc.tables:
            for row in table.rows:
                parts.append(" | ".join(c.text for c in row.cells))
        return "\n".join(parts)
    return data.decode("utf-8", errors="ignore")


def read_upload(up):
    """Return plain text from an uploaded file, showing a friendly error if it can't be read."""
    if up is None:
        return ""
    try:
        return extract_text(up.name, up.getvalue())
    except ImportError:
        st.error("Missing library. In Terminal run: python3 -m pip install pypdf python-docx")
    except Exception as e:
        st.error(f"Could not read '{up.name}': {e}")
    return ""


# ---------------------- evidence retrieval ----------------------
@st.cache_data(ttl=3600, show_spinner=False)
def epmc_search(query, n=10):
    r = requests.get(
        EPMC_URL,
        params={"query": query, "format": "json", "resultType": "core", "pageSize": n},
        timeout=20,
    )
    r.raise_for_status()
    data = r.json()
    papers = []
    for p in data.get("resultList", {}).get("result", []):
        papers.append({
            "Title": (p.get("title") or "").rstrip("."),
            "Year": p.get("pubYear"),
            "PMID": p.get("pmid"),
            "DOI": p.get("doi"),
            "Abstract": (p.get("abstractText") or "")[:1200],
        })
    return int(data.get("hitCount", 0)), papers


def extract_terms(text, limit=6):
    words = re.findall(r"[A-Za-z][A-Za-z\-]{3,}", text.lower())
    seen, terms = set(), []
    for w in words:
        if w not in STOP and w not in seen:
            seen.add(w)
            terms.append(w)
    return terms[:limit]


def _clean(t):
    return re.sub(r'["\\]', "", str(t)).strip().lower()[:60]


def simple_words(text, limit=5):
    words = re.findall(r"[A-Za-z0-9][A-Za-z0-9\-]+", text.lower())
    out = []
    for w in words:
        if w not in STOP and w not in out:
            out.append(w)
    return out[:limit]


def parse_keywords(text, is_list=False):
    """Turn a few typed keywords (or a sentence) into concept dicts. Commas keep phrases together."""
    text = text.strip()
    if is_list or len(text.split()) <= 10:
        parts = [p.strip() for p in re.split(r"[,;\n]+", text) if p.strip()]
        if len(parts) < 2:
            parts = simple_words(text)
    else:
        parts = extract_terms(text, limit=5)
    concepts = []
    for p in parts:
        name = _clean(p)
        if name and name not in [c["name"] for c in concepts]:
            concepts.append({"name": name, "synonyms": []})
    return concepts[:5]


def concept_clause(c):
    terms = [c["name"]] + c["synonyms"]
    inner = " OR ".join(f'TITLE_ABS:"{t}"' for t in terms)
    return f"({inner})" if len(terms) > 1 else inner


def build_query(concepts):
    """Title/abstract only (Europe PMC searches full text by default, which inflates hit counts)."""
    return " AND ".join(concept_clause(c) for c in concepts)


def concept_hit(text, c):
    for w in [c["name"]] + c["synonyms"]:
        if re.search(r"(?<!\w)" + re.escape(w) + r"(?:s|es)?(?!\w)", text):
            return True
    return False


def score_papers(concepts, papers):
    """Rank retrieved papers by how many of your concepts appear (title counts more than abstract)."""
    for p in papers:
        title = p["Title"].lower()
        body = (p["Title"] + " " + p["Abstract"]).lower()
        in_title = sum(concept_hit(title, c) for c in concepts)
        in_body = sum(concept_hit(body, c) for c in concepts)
        p["Match %"] = round(100 * (0.6 * in_body + 0.4 * in_title) / len(concepts))
    year = lambda p: int(p["Year"]) if str(p.get("Year") or "").isdigit() else 0
    return sorted(papers, key=lambda p: (-p["Match %"], -year(p)))


def hit_novelty(hits):
    """0 hits -> 100% ; ~10,000 hits -> 0%."""
    return max(0.0, 100.0 - 100.0 * math.log10(hits + 1) / 4.0)


# ---------------------- rule-based fallback ----------------------
def rule_scores(text, has_draft, design_name=None):
    t = text.lower()
    human = any(w in t for w in ["patient", "participant", "survey", "questionnaire",
                                 "practitioner", "doctor", "clinician", "nurse", "resident", "student", "blood", "serum", "interview", "cohort"])
    lit_only = any(w in t for w in ["scoping", "systematic review", "meta-analysis", "in silico",
                                    "bioinformatic", "database", "literature"])
    wet_lab = any(w in t for w in ["assay", "culture", "pcr", "gene", "knockout", "cell line", "animal"])
    design = any(w in t for w in ["cross-sectional", "cohort", "case-control", "scoping", "systematic",
                                  "rct", "randomi", "survey", "in vitro", "in silico", "prisma"])
    sample = any(w in t for w in ["sample", "n=", "population", "participants", "tertiary", "hospital", "centre", "center"])
    outcome = any(w in t for w in ["prevalence", "incidence", "adherence", "awareness", "score", "outcome", "level", "compliance"])
    analysis = any(w in t for w in ["regression", "chi-square", "spss", "thematic", "prisma", "anova", "t-test"])
    timeline = any(w in t for w in ["week", "month", "timeline", "phase", "milestone"])
    ethics_word = any(w in t for w in ["consent", "ethic", "iec", "irb"])

    feas = 85 - (25 if human else 0) - (20 if wet_lab else 0) + (10 if lit_only else 0)
    path = (90 if lit_only and not human else 55 if human else 70) + (10 if ethics_word else 0)
    if design_name:
        feas += {"High": 10, "Medium": 0, "Low": -25}[DESIGNS[design_name]["fit"]]
        path = (path + DESIGNS[design_name]["ethics_base"]) / 2
    meth = 15 + 25 * design + 20 * sample + 20 * outcome + 20 * analysis
    plan = 35 + 35 * timeline + (15 if has_draft else 0) + (15 if design else 0)
    clamp = lambda v: int(max(5, min(95, v)))  # rule-based estimates never claim certainty
    notes = {
        "Feasibility": "Human participants detected - needs recruitment + IEC." if human else "No participant recruitment detected.",
        "Pathway": "Ethics/regulatory path inferred from keywords only.",
        "Methodology": "Scored on whether design, population, outcome and analysis are mentioned.",
        "Work Plan": "Scored on whether a timeline/milestones are mentioned.",
    }
    return {"Feasibility": clamp(feas), "Pathway": clamp(path),
            "Methodology": clamp(meth), "Work Plan": clamp(plan)}, notes, human


# ---------------------- free LLM (Gemini) ----------------------
def _gemini_json(key, model, prompt):
    body = {"contents": [{"parts": [{"text": prompt}]}],
            "generationConfig": {"responseMimeType": "application/json", "temperature": 0.2}}
    tried = []
    for m in dict.fromkeys([model] + MODEL_FALLBACKS):
        r = requests.post(GEMINI_URL.format(model=m),
                          headers={"x-goog-api-key": key, "Content-Type": "application/json"},
                          json=body, timeout=90)
        if r.status_code == 404:  # model retired / renamed -> try the next one
            tried.append(m)
            continue
        if not r.ok:
            raise RuntimeError(f"Gemini error {r.status_code}: {r.text[:200]}")
        txt = r.json()["candidates"][0]["content"]["parts"][0]["text"]
        return json.loads(txt.replace("```json", "").replace("```", "").strip()), m
    raise RuntimeError(f"No Gemini model found (tried: {', '.join(tried)}). Set a current model name in the sidebar.")


def plan_search(key, model, text, custom):
    """AI step: work out what a few rough keywords most likely mean, and expand them into search concepts."""
    prompt = f"""You are a biomedical literature-search expert helping an Indian MBBS student.
They typed a few keywords or a rough idea. Work out what they most likely mean.
INPUT: {text[:1500]}
USER KEYWORDS: {custom}

Return ONLY JSON:
{{"interpreted_proposal": "one or two clear sentences describing the most likely research proposal",
 "study_design_guess": "",
 "concepts": [{{"name": "", "synonyms": ["", ""]}}],
 "note": "what is ambiguous or assumed"}}
Rules: 2-5 concepts ordered from most central to most specific; 2-5 synonyms each (alternate spellings,
abbreviations, MeSH-style terms, e.g. 'heart attack' -> 'myocardial infarction'); keep each term 1-4 words;
do not add concepts the user did not imply.
Avoid ambiguous words that mean something else in other fields (use 'cell transplantation' not 'implant',
'islet' not 'cell'); prefer specific biomedical phrases."""
    data, used = _gemini_json(key, model, prompt)
    concepts = []
    for c in data.get("concepts", [])[:5]:
        if not isinstance(c, dict):
            continue
        name = _clean(c.get("name", ""))
        syns = [_clean(x) for x in (c.get("synonyms") or [])[:5]]
        syns = [x for x in dict.fromkeys(syns) if x and x != name]
        if name:
            concepts.append({"name": name, "synonyms": syns})
    if not concepts:
        raise ValueError("planner returned no concepts")
    return {"interpreted": str(data.get("interpreted_proposal", ""))[:600],
            "design": str(data.get("study_design_guess", ""))[:200],
            "note": str(data.get("note", ""))[:400], "concepts": concepts}, used


def rate_papers(key, model, proposal, papers):
    """AI step: strict relevance rating so papers that merely share words are dropped."""
    items = [{"i": i, "title": p["Title"], "abstract": p["Abstract"][:300]} for i, p in enumerate(papers)]
    prompt = f"""Rate how relevant each paper is to this research proposal.
PROPOSAL: {proposal}
PAPERS: {json.dumps(items)}

Return ONLY JSON: {{"ratings": [{{"i": 0, "r": 2}}]}}
r = 2: directly addresses the same question or topic.
r = 1: related background (same field, different question).
r = 0: unrelated (different topic, even if it shares some words).
Be strict: shared words alone do not make a paper relevant."""
    data, _ = _gemini_json(key, model, prompt)
    out = {}
    for x in data.get("ratings", []):
        try:
            out[int(x["i"])] = max(0, min(2, int(x["r"])))
        except Exception:
            pass
    return out


def gemini_analyze(key, model, proposal, papers, stats, extra_text, reference=""):
    ev = [{"pmid": p["PMID"], "year": p["Year"], "title": p["Title"], "abstract": p["Abstract"]} for p in papers]
    prompt = f"""You review medical research proposals for Indian MBBS students (ICMR STS style).
Use ONLY the evidence below, and ignore any paper that is not about the same topic. Cite PMIDs from the evidence list only. Never invent papers or numbers.
If evidence is thin, say so in "limits".
If EXTRA is a draft, review it in "draft_review" and cross-check its claims against the evidence; otherwise leave draft_review empty.
If REFERENCE lists ICMR STS types of study and subject areas, fill "sts" with exactly one option from each list; otherwise leave it empty.
In "blueprint", for each of Title, Introduction, Objectives, Methodology, Implications, References list 2-4 specific points THIS proposal should cover to be complete. Do not write the section text and do not invent references.

PROPOSAL: {proposal}
SEARCH STATS: {json.dumps(stats)}
EVIDENCE: {json.dumps(ev)}
{reference}
EXTRA (guideline/draft text): {extra_text[:12000]}

Return ONLY JSON:
{{"scores": {{"Novelty": 0-100, "Feasibility": 0-100, "Pathway": 0-100, "Methodology": 0-100, "Work Plan": 0-100}},
 "notes": {{"Novelty": "", "Feasibility": "", "Pathway": "", "Methodology": "", "Work Plan": ""}},
 "already_done": ["short items with PMID"],
 "gap": "what is still unexplored",
 "methodology": {{"design": "", "population": "", "sample_size_approach": "", "tools": "", "analysis": ""}},
 "work_plan": [{{"weeks": "", "task": ""}}],
 "ethics": {{"human_participants": true, "likely_category": "exempt / expedited / full board", "reasons": ""}},
 "sts": {{"type_of_study": "", "subject_area": ""}},
 "blueprint": [{{"section": "", "cover": ["", ""]}}],
 "draft_review": {{"strengths": [], "problems": [], "contradictions": []}},
 "limits": ""}}"""
    return _gemini_json(key, model, prompt)


# ---------------------- chart ----------------------
def radar(scores):
    vals = [scores[p] for p in PILLARS]
    labels = [f"<b>{p}</b><br>{scores[p]}%" for p in PILLARS]
    grid = "rgba(127,127,127,0.35)"
    fig = go.Figure(go.Scatterpolar(
        r=vals + vals[:1], theta=labels + labels[:1], fill="toself", mode="lines+markers",
        line=dict(color="#2A7F8E", width=3), marker=dict(size=8, color="#1F3A5F"), fillcolor="rgba(42,127,142,0.30)",
    ))
    fig.update_layout(
        polar=dict(
            gridshape="linear", bgcolor="rgba(0,0,0,0)",
            radialaxis=dict(range=[0, 100], tickvals=[20, 40, 60, 80, 100], ticksuffix="%", showline=False,
                            gridcolor=grid, tickfont=dict(size=10)),
            angularaxis=dict(rotation=90, direction="clockwise", gridcolor=grid, tickfont=dict(size=13)),
        ),
        showlegend=False, height=470, margin=dict(l=70, r=70, t=30, b=30), paper_bgcolor="rgba(0,0,0,0)",
    )
    return fig


CSS = """
<style>
.block-container {padding-top: 1.2rem; max-width: 1250px;}
.hero {background: linear-gradient(120deg, #1F3A5F 0%, #2A7F8E 100%); color: #fff; padding: 26px 32px; border-radius: 18px; margin-bottom: 20px;}
.hero .t {font-size: 2.1rem; font-weight: 800; line-height: 1.15; color: #fff;}
.hero p {color: rgba(255,255,255,.9); margin: 6px 0 0 0; font-size: 1.02rem;}
.steps {display: flex; gap: 10px; margin-top: 16px; flex-wrap: wrap;}
.step {background: rgba(255,255,255,.17); padding: 5px 13px; border-radius: 999px; font-size: .83rem; color: #fff;}
.card {border: 1px solid rgba(127,127,127,.30); background: rgba(127,127,127,.07); border-radius: 14px; padding: 14px 18px;}
.card .lbl {font-size: .76rem; text-transform: uppercase; letter-spacing: .06em; opacity: .72;}
.card .val {font-size: 2.1rem; font-weight: 800; line-height: 1.15; margin-top: 2px;}
.card .val.small {font-size: 1.15rem; line-height: 1.3; margin-top: 8px;}
.card .sub {font-size: .82rem; opacity: .78; margin-top: 3px;}
.bar {height: 8px; border-radius: 6px; background: rgba(127,127,127,.25); overflow: hidden; margin-top: 8px;}
.bar > span {display: block; height: 100%; border-radius: 6px;}
.prow {margin-bottom: 14px;}
.prow .top {display: flex; justify-content: space-between; font-weight: 700; gap: 12px;}
.prow .why {font-size: .82rem; opacity: .78; margin-top: 4px;}
.section-title {font-size: 1.08rem; font-weight: 800; margin: 4px 0 10px 0;}
div[data-testid="stTabs"] button {font-weight: 600; font-size: 1rem;}
</style>
"""


def _tone(p):
    return "#1e8e5a" if p >= 70 else "#d9a400" if p >= 40 else "#c0392b"


def card_html(label, value, sub="", pct=None, small=False):
    bar = f'<div class="bar"><span style="width:{max(0, min(100, pct))}%;background:{_tone(pct)}"></span></div>' if pct is not None else ""
    return (f'<div class="card"><div class="lbl">{escape(str(label))}</div>'
            f'<div class="val{" small" if small else ""}">{escape(str(value))}</div>'
            f'<div class="sub">{escape(str(sub))}</div>{bar}</div>')


def pillar_row(name, score, why=""):
    return (f'<div class="prow"><div class="top"><span>{escape(str(name))}</span>'
            f'<span style="color:{_tone(score)}">{score}%</span></div>'
            f'<div class="bar"><span style="width:{max(0, min(100, score))}%;background:{_tone(score)}"></span></div>'
            f'<div class="why">{escape(str(why))}</div></div>')


def verdict(n):
    if n >= 75:
        return "Highly novel - confirm no unindexed duplicates"
    if n >= 50:
        return "Moderately novel - sharpen your unique angle"
    if n >= 25:
        return "Crowded topic - narrow the population, setting or outcome"
    return "Heavily studied - consider a different angle"


def build_report(proposal, scores, notes, queries, stats, ai, audit=None, design_name=None, extra_lines=None):
    lines = ["RESEARCH IDEA ANALYSIS", "", f"Proposal: {proposal}", "", "SCORES"]
    lines += [f"- {p}: {scores[p]}%  ({notes.get(p, '')})" for p in PILLARS]
    lines += ["", "SEARCHES"] + [f"- {q} -> {s['hits']} hits" for q, s in zip(queries, stats)]
    if ai:
        lines += ["", f"GAP: {ai.get('gap', '')}", f"METHODOLOGY: {json.dumps(ai.get('methodology', {}))}",
                  f"ETHICS: {json.dumps(ai.get('ethics', {}))}", f"LIMITS: {ai.get('limits', '')}"]
        lines += [f"- {w.get('weeks')}: {w.get('task')}" for w in ai.get("work_plan", [])]
    if design_name:
        d = DESIGNS[design_name]
        lines += ["", f"DESIGN GUIDE: {design_name}", f"- Guideline: {d['guideline']}", f"- Sample size: {d['sample_size']}",
                  f"- Analysis: {d['analysis']}", f"- Ethics: {d['ethics']}"]
    if audit:
        lines += ["", f"DRAFT CHECK: {audit['words']} words (limit {audit['limit'] or 'not set'})",
                  "Missing sections: " + (", ".join(k for k, v in audit["sections"].items() if not v) or "none"),
                  "Placeholders: " + (", ".join(audit["placeholders"]) or "none")]
        lines += [f"- {d['DOI']}: {d['Status']}" for d in audit["dois"]]
    lines += extra_lines or []
    lines += ["", "Advisory only. Final ethics decision rests with your IEC."]
    return "\n".join(lines)


# ---------------------- uploaded draft check ----------------------
SECTIONS = {
    "Aim / objectives": ["aim", "objective"],
    "Methodology": ["methodology", "methods", "study design"],
    "Sample size / population": ["sample size", "population", "inclusion"],
    "Ethics / consent": ["ethic", "consent", "iec", "irb"],
    "Timeline": ["timeline", "week", "month", "milestone"],
    "References": ["reference", "doi", "et al"],
}
PLACEHOLDERS = ["tbd", "to be decided", "pending", "unconfirmed", "xx", "lorem", "[insert", "todo"]


@st.cache_data(ttl=86400, show_spinner=False)
def check_doi(doi):
    try:
        r = requests.get("https://api.crossref.org/works/" + quote(doi, safe="/"),
                         headers={"User-Agent": "Apostle9/1.0"}, timeout=10)
        if r.status_code == 200:
            return "Found", ((r.json().get("message", {}).get("title") or [""])[0])
        if r.status_code == 404:
            return "Not in Crossref (typo, or a non-Crossref DOI)", ""
        return f"Check failed ({r.status_code})", ""
    except Exception:
        return "Check failed (network)", ""


STS_HEADINGS = {
    "Title": r"title",
    "Introduction": r"introduction|background",
    "Objectives": r"objectives?|aims?(?: and objectives?)?",
    "Methodology": r"methodology|materials and methods|methods",
    "Implications / outcome": r"implications?(?:\s*/\s*(?:study\s*)?outcomes?)?|expected outcomes?|study outcomes?",
    "References": r"references?",
}


def sts_spans(text):
    """Text of each STS concept-note section, found by its heading."""
    found = {}
    for name, pat in STS_HEADINGS.items():
        m = re.search(r"^[ \t]*(?:\d+[.)][ \t]*)?(?:%s)\b" % pat, text, re.I | re.M)
        if m:
            found[name] = m
    order = sorted(found.items(), key=lambda kv: kv[1].start())
    spans = {}
    for i, (name, m) in enumerate(order):
        end = order[i + 1][1].start() if i + 1 < len(order) else len(text)
        spans[name] = text[m.end():end]
    if "Title" not in spans and order:
        head = text[:order[0][1].start()]
        head = "\n".join(ln for ln in head.splitlines() if "reference id" not in ln.lower())
        if head.strip():
            spans["Title"] = head
    return spans


def sts_section_check(text):
    """Count words per section of an ICMR STS concept note against the official guide."""
    spans = sts_spans(text)
    rows = []
    for name, guide, cap in STS_FORMAT:
        if name not in spans:
            rows.append({"Section": name, "Words": "-", "Guide": guide, "Status": "⚠️ heading not found"})
            continue
        wc = len(spans[name].split())
        if wc <= guide:
            status = "✅ within guide"
        elif wc <= cap:
            status = f"⚠️ {wc - guide} over guide (up to +100 allowed)"
        else:
            status = f"❌ over the limit by {wc - cap}"
        rows.append({"Section": name, "Words": wc, "Guide": guide, "Status": status})
    return rows


def identity_flags(text):
    """Details the STS format says must not appear (such proposals are rejected)."""
    flags = []
    if re.search(r"[\w.+-]+@[\w-]+\.[\w.]+", text):
        flags.append("an email address")
    if re.search(r"(?<!\d)(?:\+?91[\s-]?)?[6-9]\d{9}(?!\d)", text):
        flags.append("a phone number")
    if re.search(r"\bDr\.?\s+[A-Z][a-z]+", text):
        flags.append("a person's name after 'Dr' (guide?)")
    if re.search(r"\b[A-Z][A-Za-z.]+(?: [A-Z][A-Za-z.]+)* (?:Hospital|Medical College|Institute|University)\b", text):
        flags.append("a named hospital/college (give only the district)")
    return flags


def audit_draft(text, word_limit, design_name=None, framework=None):
    t = text.lower()
    words = len(text.split())
    sections = {k: any(w in t for w in v) for k, v in SECTIONS.items()}
    holders = [h for h in PLACEHOLDERS if re.search(r"(?<!\w)" + re.escape(h) + r"(?!\w)", t)]
    dois = sorted({d.rstrip(".,;)]") for d in re.findall(r"10\.\d{4,9}/[^\s\"<>]+", text)})[:15]
    doi_rows = [{"DOI": d, "Status": (c := check_doi(d))[0], "Registered title": c[1]} for d in dois]
    sts = framework == "ICMR STS"
    return {"words": words, "limit": word_limit, "sections": sections, "placeholders": holders, "dois": doi_rows,
            "design": design_name, "design_checks": design_checks(design_name, text),
            "sts_sections": sts_section_check(text) if sts else [], "identity": identity_flags(text) if sts else []}


# ---------------------- reference library (old ICMR files, journal papers) ----------------------
try:
    LIB_DIR = Path(__file__).parent / "reference_library"
except NameError:
    LIB_DIR = Path.cwd() / "reference_library"
LIB_EXT = (".txt", ".md", ".csv", ".pdf", ".docx")
LIB_STOP = STOP | set("the and for with that this from were was are have has not but can may per our their its who whom".split())


@st.cache_data(show_spinner=False)
def load_bundled_library():
    docs = []
    if LIB_DIR.is_dir():
        for f in sorted(LIB_DIR.iterdir()):
            if f.is_file() and f.suffix.lower() in LIB_EXT:
                try:
                    txt = extract_text(f.name, f.read_bytes())
                except Exception:
                    continue
                if txt.strip():
                    docs.append({"source": f.name, "text": txt})
    return docs


def _tok(t):
    return [w for w in re.findall(r"[a-z0-9]{3,}", t.lower()) if w not in LIB_STOP]


def _units(text, size=60):
    """Split a document into comparable units: one line (or a few short wrapped lines) each; long paragraphs become windows."""
    units, buf = [], []
    for line in text.splitlines():
        w = line.split()
        if not w:
            continue
        buf += w
        if len(buf) >= 8:
            if len(buf) <= size * 2:
                units.append(" ".join(buf))
            else:
                for i in range(0, max(1, len(buf) - size // 2), size // 2):
                    units.append(" ".join(buf[i:i + size]))
            buf = []
    if buf:
        units.append(" ".join(buf))
    return units


def search_library(query_text, docs, top=5, max_units=8000):
    """Word-overlap (TF-IDF cosine) search of your reference files. Not semantic: it matches shared words."""
    units = []
    for d in docs:
        for u in _units(d["text"]):
            units.append((d["source"], u))
            if len(units) >= max_units:
                break
        if len(units) >= max_units:
            break
    q = _tok(query_text)
    if not units or not q:
        return []
    toks = [_tok(u[1]) for u in units]
    df = {}
    for ts in toks:
        for w in set(ts):
            df[w] = df.get(w, 0) + 1
    n_units = len(units)

    def vec(ts):
        tf = {}
        for w in ts:
            tf[w] = tf.get(w, 0) + 1
        return {w: c * (math.log((n_units + 1) / (df.get(w, 0) + 1)) + 1) for w, c in tf.items()}

    qv = vec(q)
    qn = math.sqrt(sum(v * v for v in qv.values()))
    scored = []
    for (src, seg), ts in zip(units, toks):
        v = vec(ts)
        n = math.sqrt(sum(x * x for x in v.values()))
        if n:
            scored.append((sum(qv[w] * v[w] for w in qv if w in v) / (qn * n), src, seg, set(ts)))
    scored.sort(key=lambda x: -x[0])
    out, kept = [], []
    for score, src, seg, tset in scored:
        if score <= 0 or len(out) >= top:
            break
        if any(src == k[0] and len(tset & k[1]) / max(1, len(tset | k[1])) > 0.6 for k in kept):
            continue  # near-duplicate of a result already listed
        kept.append((src, tset))
        out.append({"Source": src, "Match %": round(100 * score), "Passage": seg[:300] + ("…" if len(seg) > 300 else "")})
    return out


def journal_check(main_query, journals):
    """How many papers on your topic each chosen journal has indexed in Europe PMC."""
    rows = []
    for j in journals[:6]:
        j = j.replace('"', "").strip()
        if not j:
            continue
        try:
            hits, found = epmc_search(f'({main_query}) AND JOURNAL:"{j}"', 3)
            rows.append({"Journal": j, "Matching papers": hits,
                         "Top result": found[0]["Title"] if found else "-", "Year": found[0]["Year"] if found else ""})
        except Exception as e:
            rows.append({"Journal": j, "Matching papers": None, "Top result": f"lookup failed: {str(e)[:60]}", "Year": ""})
    return rows


# ---------------------- document completion + 100% blueprint ----------------------
BLUEPRINT_STS = {
    "Title": "Short, accurate and concise; makes the central objective and variables clear to a reviewer (25 words max).",
    "Introduction": "Rationale and significance of the work; the specific gap or controversy in current knowledge, supported by recent literature.",
    "Objectives": "Specific, achievable statements of what will be described, determined, identified, compared or confirmed; a hypothesis if relevant.",
    "Methodology": "Study type and design with justification, population, sample size and selection criteria, intervention (if any), data collection and instruments, quality control, confidentiality, analysis plan, ethical considerations (district only for the site).",
    "Implications / outcome": "What the study is expected to add scientifically, and the skills and knowledge the student will gain.",
    "References": "Recent evidence in Vancouver style, only PubMed-indexed published articles that support the proposal.",
}
BLUEPRINT_GENERIC = {
    "Title": "Short, accurate and informative.",
    "Introduction": "Background, rationale and the gap in knowledge, with recent literature.",
    "Objectives": "Clear primary and secondary objectives or a hypothesis.",
    "Methodology": "Design, setting, population, sample size, data collection, analysis plan and ethics.",
    "Implications / outcome": "Expected outcomes and significance.",
    "References": "Recent, verifiable references in a consistent style.",
}


def completion_level(p):
    if p < 40:
        return "Early draft"
    if p < 70:
        return "Developing"
    if p < 90:
        return "Nearly complete"
    return "Complete structure - ready for guide/expert review"


def compute_completion(audit, text, framework, design_name):
    """Percent completeness of an uploaded draft against the format and checklist (structure, not scientific quality)."""
    if not audit:
        return None
    comps, gaps = [], []

    def add(name, got, mx, note=""):
        comps.append({"Component": name, "Score": round(got, 1), "Max": round(mx, 1), "Note": note})

    if framework == "ICMR STS" and audit.get("sts_sections"):
        weights = {"Title": 5, "Introduction": 10, "Objectives": 8, "Methodology": 20, "Implications / outcome": 5, "References": 7}
        caps = {n: c for n, _, c in STS_FORMAT}
        for row in audit["sts_sections"]:
            name, words, guide = row["Section"], row["Words"], row["Guide"]
            w = weights[name]
            if words == "-":
                add(name, 0, w, "missing")
                gaps.append(f"{name}: add this section (about {guide} words).")
                continue
            if name == "Title":
                frac = 1.0 if 3 <= words <= 25 else (0.5 if words > 25 else 0.3)
                if words > 25:
                    gaps.append("Title: shorten to 25 words or fewer.")
                elif words < 3:
                    gaps.append("Title: write a descriptive title.")
            elif name == "References":
                frac = 1.0 if words <= caps[name] else 0.7  # quantity is scored under "Reference list" below
                if words > caps[name]:
                    gaps.append(f"References: {words} words is over the limit; trim to {guide + 100} or fewer.")
            else:
                target = 0.6 * guide
                frac = min(1.0, words / target)
                if words > caps[name]:
                    frac = 0.7
                    gaps.append(f"{name}: {words} words is over the limit; trim to {guide + 100} or fewer.")
                elif frac < 1:
                    gaps.append(f"{name}: {words} words, aim for at least {int(target)} (guide {guide}).")
            add(name, w * frac, w, f"{words} words")
    else:
        generic = audit["sections"]
        w = 55 / max(1, len(generic))
        for k, v in generic.items():
            add(k, w if v else 0, w, "found" if v else "missing")
            if not v:
                gaps.append(f"Cover: {k.lower()}.")

    dc = audit.get("design_checks") or {}
    if dc:
        found = sum(dc.values())
        add("Design-specific elements", 20 * found / len(dc), 20, f"{found}/{len(dc)} present")
        for k, v in dc.items():
            if not v:
                gaps.append(f"Methodology: mention {k.lower()}.")

    eth = audit["sections"].get("Ethics / consent", False)
    add("Ethics / consent", 5 if eth else 0, 5, "mentioned" if eth else "missing")
    if not eth:
        gaps.append("Add ethics: IEC approval status and informed consent (or why exempt).")

    refs_text = sts_spans(text).get("References", "") if framework == "ICMR STS" else ""
    n_refs = len(re.findall(r"(?m)^\s*\[?\(?\d{1,3}[.)\]]\s+\S", refs_text)) if refs_text else 0
    n_refs = max(n_refs, len(audit["dois"]), text.lower().count("et al"))
    bad = sum(1 for d in audit["dois"] if str(d["Status"]).startswith("Not in Crossref"))
    frac = max(0.0, min(1.0, n_refs / 5) - 0.2 * bad)
    add("Reference list (quantity, DOIs)", 10 * frac, 10, f"{n_refs} detected, {bad} DOI(s) not found")
    if n_refs < 5:
        gaps.append(f"References: add at least {5 - n_refs} more recent, PubMed-indexed reference(s).")
    if bad:
        gaps.append("References: fix or remove DOI(s) not found in Crossref.")

    penalty = 2.5 * len(audit.get("identity", [])) + 2 * len(audit["placeholders"])
    add("Clean text (no identifiers or placeholders)", max(0, 10 - penalty), 10,
        "clean" if penalty == 0 else "needs cleaning")
    if audit.get("identity"):
        gaps.append("Remove identifying details: " + ", ".join(audit["identity"]) + ".")
    if audit["placeholders"]:
        gaps.append("Replace placeholder text: " + ", ".join(audit["placeholders"]) + ".")

    got, mx = sum(c["Score"] for c in comps), sum(c["Max"] for c in comps)
    pct = int(round(100 * got / mx)) if mx else 0
    if gaps:  # never call a draft complete while open items remain
        pct = min(pct, 99)
    level = completion_level(min(pct, 89) if gaps else pct)
    return {"percent": pct, "level": level, "components": comps, "gaps": gaps}


def build_blueprint(audit, design_name, framework, ai, papers):
    """What each section of a complete (100%) document should cover, tailored to this idea."""
    base = BLUEPRINT_STS if framework == "ICMR STS" else BLUEPRINT_GENERIC
    targets = {n: f"{g} words" for n, g, _ in STS_FORMAT} if framework == "ICMR STS" else {}
    now = {}
    if audit and audit.get("sts_sections"):
        now = {r["Section"]: (f"{r['Words']} words" if r["Words"] != "-" else "Missing") for r in audit["sts_sections"]}
    ai_points = {}
    for b in (ai or {}).get("blueprint", []) or []:
        if isinstance(b, dict) and b.get("section"):
            ai_points[str(b["section"]).lower().split()[0].strip("/,:")] = [str(x) for x in (b.get("cover") or [])][:4]
    tailored = {}
    if design_name:
        d = DESIGNS[design_name]
        tailored["Methodology"] = [f"Design: {design_name} (report per {d['guideline']})", f"Sample size: {d['sample_size']}",
                                   f"Analysis: {d['analysis']}"]
    if ai and ai.get("gap"):
        tailored["Introduction"] = [f"State the gap: {ai['gap']}"]
    pmids = [p["PMID"] for p in (papers or [])[:3] if p.get("PMID")]
    if pmids:
        tailored["References"] = ["Start from the closest relevant papers found (PMIDs " + ", ".join(pmids) + ") and add recent ones"]
    rows = []
    for name, desc in base.items():
        pts = ai_points.get(name.lower().split()[0].strip("/,:"), []) + tailored.get(name, [])
        rows.append({"Section": name, "Target": targets.get(name, "-"), "A complete section covers": desc,
                     "Your draft now": now.get(name, "-" if audit else "Not started"),
                     "Specific to your idea": "; ".join(pts) if pts else "-"})
    return rows


# ---------------------- PDF report ----------------------
def pdf_text(s):
    s = "" if s is None else str(s)
    for a, b in (("✅", "[OK]"), ("⚠️", "[!]"), ("⚠", "[!]"), ("❌", "[X]"), ("…", "..."), ("₹", "Rs "), ("–", "-"), ("—", "-"),
                 ("‘", "'"), ("’", "'"), ("“", '"'), ("”", '"'), ("≥", ">="), ("≤", "<="), ("→", "->"), ("\u200b", ""), ("🧠", "")):
        s = s.replace(a, b)
    return s.encode("cp1252", "replace").decode("cp1252")


def build_pdf(res, framework):
    """Official-style PDF: summary, pentagon, novelty, completion, checks, evidence, AI assessment."""
    import datetime
    import hashlib
    from reportlab.graphics.shapes import Circle, Drawing, Line, Polygon, Rect, String
    from reportlab.lib import colors
    from reportlab.lib.pagesizes import A4
    from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
    from reportlab.lib.units import mm
    from reportlab.platypus import KeepTogether, Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle

    NAVY, TEAL, LIGHT, GREY = colors.HexColor("#1F3A5F"), colors.HexColor("#2A7F8E"), colors.HexColor("#EEF2F7"), colors.HexColor("#8A94A6")
    W = 174 * mm
    base = getSampleStyleSheet()
    body = ParagraphStyle("body", parent=base["BodyText"], fontName="Helvetica", fontSize=9.5, leading=13)
    small = ParagraphStyle("small", parent=body, fontSize=8, leading=10.5)
    smallb = ParagraphStyle("smallb", parent=small, fontName="Helvetica-Bold")
    hdr = ParagraphStyle("hdr", parent=smallb, textColor=colors.white)
    h1 = ParagraphStyle("h1", parent=body, fontName="Helvetica-Bold", fontSize=13, textColor=NAVY, spaceBefore=12, spaceAfter=5)
    mono = ParagraphStyle("mono", parent=small, fontName="Courier", fontSize=7.5, leading=9.5)

    def P(t, st=body):
        return Paragraph(escape(pdf_text(t)), st)

    def tone(p):
        return TEAL if p >= 70 else colors.HexColor("#D9A400") if p >= 40 else colors.HexColor("#C0392B")

    def bar(pct, w=150, h=9):
        pct = max(0, min(100, pct))
        d = Drawing(w + 34, h + 3)
        d.add(Rect(0, 1, w, h, fillColor=colors.HexColor("#E3E8EF"), strokeColor=None))
        d.add(Rect(0, 1, w * pct / 100, h, fillColor=tone(pct), strokeColor=None))
        d.add(String(w + 5, 2.5, f"{int(pct)}%", fontName="Helvetica-Bold", fontSize=8, fillColor=NAVY))
        return d

    def pentagon(scores, size=290):
        import math
        d = Drawing(size, size)
        cx = cy = size / 2
        R = size * 0.31
        pt = lambda i, f: (cx + R * f * math.cos(math.pi / 2 - 2 * math.pi * i / 5), cy + R * f * math.sin(math.pi / 2 - 2 * math.pi * i / 5))
        for lvl in (0.2, 0.4, 0.6, 0.8, 1.0):
            d.add(Polygon([c for i in range(5) for c in pt(i, lvl)], strokeColor=colors.HexColor("#C9D1DC"), fillColor=None, strokeWidth=0.6))
            x, y = pt(0, lvl)
            d.add(String(x + 3, y - 2, f"{int(lvl * 100)}", fontName="Helvetica", fontSize=6, fillColor=GREY))
        for i in range(5):
            x, y = pt(i, 1)
            d.add(Line(cx, cy, x, y, strokeColor=colors.HexColor("#C9D1DC"), strokeWidth=0.6))
        vals = [scores[n] / 100 for n in PILLARS]
        d.add(Polygon([c for i, v in enumerate(vals) for c in pt(i, v)], strokeColor=NAVY, fillColor=TEAL, fillOpacity=0.3, strokeWidth=1.6))
        for i, v in enumerate(vals):
            x, y = pt(i, v)
            d.add(Circle(x, y, 2.6, fillColor=NAVY, strokeColor=NAVY))
        for i, n in enumerate(PILLARS):
            x, y = pt(i, 1.2)
            d.add(String(x, y + 3, n, fontName="Helvetica-Bold", fontSize=9, fillColor=NAVY, textAnchor="middle"))
            d.add(String(x, y - 8, f"{scores[n]}%", fontName="Helvetica", fontSize=9, fillColor=colors.black, textAnchor="middle"))
        d.hAlign = "CENTER"
        return d

    def table(rows, widths, header=True):
        data = [[(P(c, hdr if (header and i == 0) else small) if isinstance(c, str) else c) for c in r] for i, r in enumerate(rows)]
        t = Table(data, colWidths=widths, repeatRows=1 if header else 0)
        cmds = [("VALIGN", (0, 0), (-1, -1), "TOP"), ("GRID", (0, 0), (-1, -1), 0.4, colors.HexColor("#C9D1DC")),
                ("LEFTPADDING", (0, 0), (-1, -1), 4), ("RIGHTPADDING", (0, 0), (-1, -1), 4),
                ("TOPPADDING", (0, 0), (-1, -1), 3), ("BOTTOMPADDING", (0, 0), (-1, -1), 3)]
        if header:
            cmds += [("BACKGROUND", (0, 0), (-1, 0), NAVY), ("ROWBACKGROUNDS", (0, 1), (-1, -1), [colors.white, LIGHT])]
        t.setStyle(TableStyle(cmds))
        return t

    story, counter = [], [0]

    def H(title):
        counter[0] += 1
        story.append(P(f"{counter[0]}. {title}", h1))

    def bullets(items):
        for it in items:
            story.append(P("• " + str(it), body))

    scores, comp, ai, audit = res["scores"], res.get("completion"), res.get("ai"), res.get("audit")
    today = datetime.date.today()
    rid = "RIA-" + today.strftime("%Y%m%d") + "-" + hashlib.sha1(res["proposal"].encode("utf-8", "ignore")).hexdigest()[:6].upper()

    # ---- title block ----
    title_style = ParagraphStyle("t", parent=body, fontName="Helvetica-Bold", fontSize=17, textColor=colors.white, leading=21)
    sub_style = ParagraphStyle("t2", parent=small, textColor=colors.white, fontSize=9)
    head = Table([[P("RESEARCH IDEA ASSESSMENT REPORT", title_style)],
                  [P(f"Prepared for: {framework}   |   Date: {today.strftime('%d %B %Y')}   |   Report ID: {rid}", sub_style)]], colWidths=[W])
    head.setStyle(TableStyle([("BACKGROUND", (0, 0), (-1, -1), NAVY), ("LEFTPADDING", (0, 0), (-1, -1), 10),
                              ("TOPPADDING", (0, 0), (-1, 0), 12), ("BOTTOMPADDING", (0, -1), (-1, -1), 10)]))
    story += [head, Spacer(1, 8), P("Proposal under review", smallb), P(res["proposal"])]

    # ---- 1 summary ----
    H("Summary")
    overall = round(sum(scores.values()) / len(scores))
    rows = [["Indicator", "Result"],
            ["Novelty (literature-based)", bar(scores["Novelty"])],
            ["Pentagon average (5 pillars)", bar(overall)],
            ["Document completion", bar(comp["percent"]) if comp else "No draft uploaded"],
            ["Detected study design", res.get("design") or "Not detected"],
            ["Novelty reading", verdict(scores["Novelty"])]]
    if comp:
        rows.append(["Document level", comp["level"]])
    story.append(table(rows, [55 * mm, W - 55 * mm]))

    # ---- 2 pentagon ----
    H("Pentagon assessment")
    story.append(pentagon(scores))
    story.append(table([["Pillar", "Score", "Basis"]] + [[n, f"{scores[n]}%", res["notes"].get(n, "")] for n in PILLARS],
                       [28 * mm, 16 * mm, W - 44 * mm]))
    story.append(P(res["source"], small))

    # ---- 3 novelty ----
    H("Novelty analysis")
    story.append(table([["Search (broad to specific)", "Papers found", "Novelty at this stage"]] +
                       [[s["query"], f"{s['hits']:,}", bar(s["novelty"], 100)] for s in res["stats"]],
                       [W - 80 * mm, 24 * mm, 56 * mm]))
    story += [Spacer(1, 4), P("Europe PMC query behind the novelty score:", smallb), P(res["queries"][res["mid"]], mono)]
    plan = res.get("plan") or {}
    if plan.get("interpreted") or plan.get("design"):
        story += [Spacer(1, 4), P(f"Interpretation of your idea: {plan.get('interpreted', '')}  Likely design: {plan.get('design', '')}", small)]
    story.append(Spacer(1, 4))
    story.append(table([["Concept", "Also searched as"]] +
                       [[c["name"], ", ".join(c["synonyms"]) or "-"] for c in res["concepts"]], [50 * mm, W - 50 * mm]))

    # ---- 4 design guide ----
    if res.get("design"):
        d = DESIGNS[res["design"]]
        H("Design guide: " + res["design"])
        story.append(table([["Item", "Guidance"], ["Reporting guideline", d["guideline"]], ["Sample size", d["sample_size"]],
                            ["Typical analysis", d["analysis"]], ["Ethics route", d["ethics"]], ["Fit for a short student project", d["fit"]]],
                           [45 * mm, W - 45 * mm]))
        story.append(P("Curated general reference; confirm details with your guide and the current call document.", small))

    # ---- 5 completion + blueprint ----
    H("Document completion and the 100% blueprint")
    if comp:
        story += [bar(comp["percent"], 220, 12), P(f"{comp['percent']}% - {comp['level']}. Completion measures structure and "
                  "completeness against the format and checklist, not scientific quality.", small), Spacer(1, 4)]
        story.append(table([["Component", "Score", "Max", "Note"]] +
                           [[c["Component"], str(c["Score"]), str(c["Max"]), c["Note"]] for c in comp["components"]],
                           [70 * mm, 18 * mm, 18 * mm, W - 106 * mm]))
        if comp["gaps"]:
            story += [Spacer(1, 4), P("To reach 100%:", smallb)]
            bullets(comp["gaps"])
    else:
        story.append(P("No draft was uploaded, so completion could not be measured. The blueprint below shows what a complete document contains."))
    story += [Spacer(1, 6), P("What the 100% document looks like", smallb)]
    story.append(table([["Section", "Target", "A complete section covers", "Your draft now", "Specific to your idea"]] +
                       [[r["Section"], r["Target"], r["A complete section covers"], r["Your draft now"], r["Specific to your idea"]] for r in res["blueprint"]],
                       [24 * mm, 16 * mm, 52 * mm, 20 * mm, W - 112 * mm]))

    # ---- 6 draft check ----
    if audit:
        H("Uploaded draft check")
        story.append(P(f"Word count: {audit['words']:,}" + (f" (limit {audit['limit']:,})" if audit["limit"] else "")))
        if audit.get("sts_sections"):
            story.append(table([["Section", "Words", "Guide", "Status"]] +
                               [[r["Section"], str(r["Words"]), str(r["Guide"]), r["Status"]] for r in audit["sts_sections"]],
                               [45 * mm, 20 * mm, 20 * mm, W - 85 * mm]))
            story.append(P("Official format: guide counts may be exceeded by up to 100 words (title 25 max). Font, spacing and "
                           "Vancouver/PubMed-only references are not checked here.", small))
        if audit.get("identity"):
            story.append(P("Identifying details found (the format says such proposals are rejected): " + ", ".join(audit["identity"]), body))
        if audit.get("design_checks"):
            story += [Spacer(1, 4), table([["Design-specific element", "In draft"]] +
                      [[k, "Yes" if v else "MISSING"] for k, v in audit["design_checks"].items()], [W - 30 * mm, 30 * mm])]
        story += [Spacer(1, 4), table([["Section", "In draft"]] + [[k, "Yes" if v else "MISSING"] for k, v in audit["sections"].items()],
                                      [W - 30 * mm, 30 * mm])]
        if audit["placeholders"]:
            story.append(P("Placeholder text found: " + ", ".join(audit["placeholders"])))
        if audit["dois"]:
            story += [Spacer(1, 4), table([["DOI", "Status", "Registered title"]] +
                      [[d["DOI"], d["Status"], d["Registered title"]] for d in audit["dois"]], [52 * mm, 42 * mm, W - 94 * mm])]

    # ---- 7 evidence ----
    H("Closest relevant papers")
    if res.get("fallback"):
        story.append(P("No paper matched all concepts; these are nearest neighbours from a broader search, not duplicates.", small))
    if res["papers"]:
        story.append(table([["Title", "Year", "Relevance", "Match %", "PMID / DOI"]] +
                           [[p["Title"], str(p.get("Year") or ""), p.get("Relevance", "-"), str(p.get("Match %", "")),
                             ("PMID " + str(p["PMID"]) if p.get("PMID") else "") + (" doi:" + p["DOI"] if p.get("DOI") else "")]
                            for p in res["papers"]], [78 * mm, 12 * mm, 20 * mm, 16 * mm, W - 126 * mm]))
    else:
        story.append(P("No closely related papers were found."))

    if res.get("journals"):
        H("Journal check")
        story.append(table([["Journal", "Matching papers", "Top result", "Year"]] +
                           [[j["Journal"], "-" if j["Matching papers"] is None else str(j["Matching papers"]), j["Top result"], str(j["Year"])]
                            for j in res["journals"]], [52 * mm, 24 * mm, W - 90 * mm, 14 * mm]))
    if res.get("library"):
        H("Reference library comparison")
        story.append(table([["Source file", "Match %", "Passage"]] +
                           [[h["Source"], str(h["Match %"]), h["Passage"]] for h in res["library"]], [40 * mm, 16 * mm, W - 56 * mm]))
        story.append(P("Matching is by shared words, not meaning; read the passages to judge real overlap.", small))

    # ---- AI assessment ----
    if ai:
        H("AI assessment")
        if ai.get("gap"):
            story.append(P("Unexplored gap: " + str(ai["gap"])))
        if ai.get("already_done"):
            story.append(P("Already done in the literature:", smallb))
            bullets(ai["already_done"])
        m = ai.get("methodology") or {}
        if m:
            story += [Spacer(1, 4), table([["Methodology", "Suggestion"]] + [[k.replace("_", " ").title(), str(v)] for k, v in m.items()],
                                          [40 * mm, W - 40 * mm])]
        if ai.get("work_plan"):
            story += [Spacer(1, 4), table([["Weeks", "Task"]] + [[str(w.get("weeks", "")), str(w.get("task", ""))] for w in ai["work_plan"]],
                                          [30 * mm, W - 30 * mm])]
        e = ai.get("ethics") or {}
        if e:
            story += [Spacer(1, 4), P("Ethics screen (advisory; the IEC decides): " + "; ".join(f"{k.replace('_', ' ')}: {v}" for k, v in e.items()))]
        sts = ai.get("sts") or {}
        if sts.get("type_of_study") or sts.get("subject_area"):
            story.append(P(f"Suggested STS form entries: type of study - {sts.get('type_of_study', '-')}; subject area - {sts.get('subject_area', '-')}"))
        dr = ai.get("draft_review") or {}
        for key_, label in (("strengths", "Strengths"), ("problems", "Problems"), ("contradictions", "Contradictions with the literature")):
            if dr.get(key_):
                story.append(P(label + ":", smallb))
                bullets(dr[key_])
        if ai.get("limits"):
            story.append(P("Limits noted by the AI: " + str(ai["limits"]), small))

    # ---- basis ----
    H("Basis, sources and limitations")
    bullets([
        "Literature evidence: Europe PMC (title/abstract search). The exact query is shown in the novelty section.",
        "Assessment engine: " + res["source"] + ".",
        "ICMR STS rules and format: DHR documents (schemes.dhr.gov.in) read on 22 October 2026; confirm on sts.icmr.org.in before submitting.",
        "Novelty reflects literature volume and the closest relevant paper; it depends on the search terms and is not a guarantee of originality.",
        "Completion measures structure and completeness, not scientific quality.",
        "This report is advisory. Ethics categories and approvals are decided only by your Institutional Ethics Committee.",
    ])

    def deco(canvas, doc):
        canvas.saveState()
        canvas.setStrokeColor(NAVY)
        canvas.setLineWidth(0.8)
        canvas.line(18 * mm, A4[1] - 14 * mm, A4[0] - 18 * mm, A4[1] - 14 * mm)
        canvas.setFont("Helvetica", 7.5)
        canvas.setFillColor(GREY)
        canvas.drawString(18 * mm, A4[1] - 11.5 * mm, "Apostle9 - advisory assessment")
        canvas.drawRightString(A4[0] - 18 * mm, A4[1] - 11.5 * mm, rid)
        canvas.line(18 * mm, 14 * mm, A4[0] - 18 * mm, 14 * mm)
        canvas.drawString(18 * mm, 9.5 * mm, "Advisory only - not an ethics, funding or scientific approval. Confirm requirements with your guide.")
        canvas.drawRightString(A4[0] - 18 * mm, 9.5 * mm, f"Page {doc.page}")
        canvas.restoreState()

    buf = io.BytesIO()
    doc = SimpleDocTemplate(buf, pagesize=A4, leftMargin=18 * mm, rightMargin=18 * mm, topMargin=20 * mm, bottomMargin=20 * mm,
                            title="Research Idea Assessment Report", author="Apostle9")
    doc.build(story, onFirstPage=deco, onLaterPages=deco)
    return buf.getvalue()


def safe_pdf(res, framework):
    try:
        return build_pdf(res, framework), None
    except ImportError:
        return None, "PDF needs reportlab. In Terminal run: python3 -m pip install reportlab"
    except Exception as e:
        return None, f"PDF could not be built: {e}"


# ---------------------- analysis ----------------------
def run_analysis(proposal, custom, key, model, extra, word_limit=0, extra_rules="",
                 framework="ICMR STS", journals=(), library_docs=None):
    plan = {"interpreted": "", "design": "", "note": "", "error": None, "ai": False}
    concepts = None
    if key:
        try:
            with st.spinner("Interpreting your idea..."):
                planned, _ = plan_search(key, model, proposal, custom)
            concepts = planned.pop("concepts")
            plan.update(planned, ai=True)
        except Exception as e:
            plan["error"] = str(e)
    if not concepts:
        concepts = parse_keywords(custom, is_list=True) if custom.strip() else parse_keywords(proposal)
    if not concepts:
        raise ValueError("Could not find any keywords - type a few words describing your idea.")
    if plan["ai"] and plan["interpreted"] and len(proposal.split()) < 25:
        proposal = plan["interpreted"]  # downstream scoring works on the clearer wording

    n = len(concepts)
    ks = sorted({min(2, n), min(4, n), n})
    queries = [build_query(concepts[:k]) for k in ks]
    labels = [" AND ".join(c["name"] for c in concepts[:k]) for k in ks]
    mid = len(ks) // 2

    stats, pool, fallback = [], [], False
    with st.status("Searching Europe PMC...", expanded=True) as status:
        for i, q in enumerate(queries):
            hits, found = epmc_search(q, 25 if i == mid else 1)
            stats.append({"query": labels[i], "hits": hits, "novelty": round(hit_novelty(hits))})
            if i == mid:
                pool = found
            status.write(f"Search {i + 1}: `{labels[i]}` -> **{hits:,}** papers")
        if not pool:
            fallback = True  # nothing matched every concept: fetch nearest neighbours from the broad query
            _, pool = epmc_search(queries[0], 25)
        status.update(label="Search complete", state="complete", expanded=False)

    ranked = score_papers(concepts, pool)[:15]
    rated, rating_error = False, None
    if key and ranked:
        try:
            with st.spinner("Checking which papers are really relevant..."):
                ratings = rate_papers(key, model, proposal, ranked)
            if ratings:
                rated = True
                names = {2: "Direct", 1: "Related", 0: "Unrelated"}
                for i, p in enumerate(ranked):
                    p["Relevance"] = names.get(ratings.get(i), "Related")
        except Exception as e:
            rating_error = str(e)
    if rated:
        order = {"Direct": 0, "Related": 1, "Unrelated": 2}
        papers = sorted([p for p in ranked if p["Relevance"] != "Unrelated"],
                        key=lambda p: (order[p["Relevance"]], -p["Match %"]))
    else:
        papers = [p for p in ranked if p["Match %"] >= 50]
    hidden = len(ranked) - len(papers)
    sim_pool = [p for p in papers if p.get("Relevance") == "Direct"] if rated else papers
    sim = max([p["Match %"] for p in sim_pool] or [0])
    main_stat = stats[mid]
    novelty = int(round(0.5 * main_stat["novelty"] + 0.5 * (100 - sim)))

    design_name = detect_design(proposal + " " + extra[:3000], plan.get("design", ""))
    reference = build_reference(design_name, extra_rules, framework)
    rule, rule_notes, _ = rule_scores(proposal + " " + extra, bool(extra), design_name)
    scores = {"Novelty": novelty, **rule}
    notes = {"Novelty": f"{main_stat['hits']:,} title/abstract matches for '{main_stat['query']}'; "
                        f"closest {'AI-confirmed ' if rated else ''}relevant paper covers {sim}% of your concepts.", **rule_notes}
    source, ai, ai_error = "Rule-based estimates (add a Gemini key for AI assessment)", None, None

    if key:
        try:
            with st.spinner("AI assessing the evidence..."):
                ai, used = gemini_analyze(key, model, proposal, papers[:10], stats, extra, reference)
            new_scores = {p: int(max(0, min(100, ai["scores"][p]))) for p in PILLARS[1:]}
            new_notes = {p: ai.get("notes", {}).get(p, "") for p in PILLARS[1:]}
            scores.update(new_scores)
            notes.update(new_notes)
            scores["Novelty"] = int(round(0.7 * novelty + 0.3 * ai["scores"]["Novelty"]))
            notes["Novelty"] += " " + ai.get("notes", {}).get("Novelty", "")
            source = f"AI assessment ({used}) grounded on retrieved papers"
        except Exception as e:
            ai, ai_error = None, str(e)

    audit = None
    if extra.strip():
        with st.spinner("Checking your uploaded draft..."):
            audit = audit_draft(extra, word_limit, design_name, framework)

    completion = compute_completion(audit, extra, framework, design_name) if audit else None
    blueprint = build_blueprint(audit, design_name, framework, ai, papers)
    journal_rows = []
    if journals:
        with st.spinner("Comparing with journals..."):
            journal_rows = journal_check(queries[mid], list(journals))
    library_hits = []
    if library_docs:
        with st.spinner("Comparing with your reference library..."):
            terms_all = [c["name"] for c in concepts] + [x for c in concepts for x in c["synonyms"]]
            library_hits = search_library(proposal + " " + " ".join(terms_all), library_docs)

    extra_lines = []
    if completion:
        extra_lines += ["", f"DOCUMENT COMPLETION: {completion['percent']}% ({completion['level']})"] + ["- To reach 100%: " + g for g in completion["gaps"]]
    if audit and audit.get("sts_sections"):
        extra_lines += ["", "STS FORMAT CHECK"] + [f"- {r['Section']}: {r['Words']} words (guide {r['Guide']}) {r['Status']}" for r in audit["sts_sections"]]
        if audit.get("identity"):
            extra_lines += ["Identifying details found: " + ", ".join(audit["identity"])]
    if journal_rows:
        extra_lines += ["", "JOURNAL CHECK"] + [f"- {r['Journal']}: {r['Matching papers']} matching papers" for r in journal_rows]
    if library_hits:
        extra_lines += ["", "REFERENCE LIBRARY MATCHES"] + [f"- {h['Source']} ({h['Match %']}%): {h['Passage'][:120]}" for h in library_hits]

    result = {"proposal": proposal, "framework": framework, "completion": completion, "blueprint": blueprint,
              "scores": scores, "audit": audit, "plan": plan,
            "journals": journal_rows, "library": library_hits, "lib_count": len(library_docs or []),
            "concepts": concepts, "queries": queries, "mid": mid, "notes": notes, "source": source, "ai": ai,
            "ai_error": ai_error, "stats": stats, "papers": papers, "labels": labels,
            "design": design_name, "hidden": hidden, "fallback": fallback, "rated": rated, "rating_error": rating_error,
            "report": build_report(proposal, scores, notes, labels, stats, ai, audit, design_name, extra_lines)}
    result["pdf"], result["pdf_error"] = safe_pdf(result, framework)
    return result


def render(res, framework):
    framework = res.get("framework", framework)
    scores, notes, ai = res["scores"], res["notes"], res["ai"]
    comp, audit, plan = res.get("completion"), res.get("audit"), res.get("plan") or {}
    overall = round(sum(scores.values()) / len(scores))
    t_over, t_nov, t_doc, t_ai, t_rep = st.tabs(["📊 Overview", "🔎 Novelty & evidence", "📁 Document", "🧠 AI assessment", "📄 Report"])

    # ---------------- Overview ----------------
    with t_over:
        st.markdown(f'<div class="card"><div class="lbl">Proposal under review</div>'
                    f'<div style="font-size:1.05rem;margin-top:4px">{escape(res["proposal"])}</div></div>', unsafe_allow_html=True)
        st.write("")
        c1, c2, c3, c4 = st.columns(4)
        c1.markdown(card_html("Novelty", f"{scores['Novelty']}%", verdict(scores["Novelty"]), scores["Novelty"]), unsafe_allow_html=True)
        c2.markdown(card_html("Pentagon average", f"{overall}%", "Average of the 5 pillars", overall), unsafe_allow_html=True)
        if comp:
            c3.markdown(card_html("Document completion", f"{comp['percent']}%", comp["level"], comp["percent"]), unsafe_allow_html=True)
        else:
            c3.markdown(card_html("Document completion", "-", "Upload a draft to measure it"), unsafe_allow_html=True)
        c4.markdown(card_html("Study design", res.get("design") or "Not detected", "Details in the AI assessment tab", small=True),
                    unsafe_allow_html=True)
        st.write("")
        left, right = st.columns([1.15, 1], gap="large")
        with left:
            st.markdown('<div class="section-title">Pentagon assessment</div>', unsafe_allow_html=True)
            st.plotly_chart(radar(scores))
        with right:
            st.markdown('<div class="section-title">Pillar breakdown</div>', unsafe_allow_html=True)
            st.markdown("".join(pillar_row(n, scores[n], notes.get(n, "")) for n in PILLARS), unsafe_allow_html=True)
            st.caption(res["source"])
            if res.get("ai_error"):
                st.warning(f"AI step failed, showing rule-based scores. ({res['ai_error']})")
        with st.expander("🧠 How I understood your idea"):
            if plan.get("error"):
                st.warning(f"AI interpretation failed, used your words directly. ({plan['error']})")
            elif not plan.get("ai"):
                st.caption("No Gemini key: keywords were used as typed, without synonym expansion.")
            if plan.get("interpreted"):
                st.markdown(f"**Interpreted proposal:** {plan['interpreted']}")
            if plan.get("design"):
                st.markdown(f"**Likely design:** {plan['design']}")
            if plan.get("note"):
                st.caption(plan["note"])
            st.dataframe(pd.DataFrame([{"Concept": c["name"], "Also searched as": ", ".join(c["synonyms"]) or "-"}
                                       for c in res["concepts"]]), hide_index=True)
            st.caption("Europe PMC query behind the novelty score:")
            st.code(res["queries"][res["mid"]], language=None)

    # ---------------- Novelty & evidence ----------------
    with t_nov:
        a, b = st.columns([1, 1.25], gap="large")
        with a:
            st.markdown('<div class="section-title">Novelty progress (broad to specific)</div>', unsafe_allow_html=True)
            st.markdown("".join(pillar_row(x["query"], x["novelty"], f"{x['hits']:,} papers found") for x in res["stats"]),
                        unsafe_allow_html=True)
            st.markdown(card_html("Overall novelty", f"{scores['Novelty']}%", verdict(scores["Novelty"]), scores["Novelty"]),
                        unsafe_allow_html=True)
        with b:
            st.markdown('<div class="section-title">Closest relevant papers</div>', unsafe_allow_html=True)
            if res.get("fallback"):
                st.caption("No paper matched all your concepts, so these come from a broader search. "
                           "Treat them as nearest neighbours, not duplicates.")
            if res.get("rating_error"):
                st.caption(f"AI relevance check failed, filtered by keyword match only. ({res['rating_error']})")
            elif not res.get("rated"):
                st.caption("Filtered by keyword match only. Add a Gemini key for AI relevance checking.")
            if res.get("hidden"):
                st.caption(f"Hid {res['hidden']} results that looked unrelated.")
            dfp = pd.DataFrame(res["papers"])
            if not dfp.empty:
                dfp["Link"] = dfp["DOI"].apply(lambda d: f"https://doi.org/{d}" if d else "")
                cols = [c for c in ["Title", "Relevance", "Match %", "Year", "PMID", "Link"] if c in dfp.columns]
                st.dataframe(dfp[cols], hide_index=True)
            else:
                st.write("No closely related papers found - a good sign for novelty, but check your keywords.")
        if res.get("journals") or res.get("lib_count"):
            st.divider()
            j, l = st.columns(2, gap="large")
            with j:
                if res.get("journals"):
                    st.markdown('<div class="section-title">📰 Journal check</div>', unsafe_allow_html=True)
                    st.caption("Papers on your topic indexed in Europe PMC for each journal (depends on how the journal title is recorded).")
                    st.dataframe(pd.DataFrame(res["journals"]), hide_index=True)
            with l:
                if res.get("lib_count"):
                    st.markdown('<div class="section-title">🗂️ Reference library</div>', unsafe_allow_html=True)
                    st.caption(f"Compared with {res['lib_count']} file(s) by shared words, not meaning.")
                    if res.get("library"):
                        top = res["library"][0]["Match %"]
                        st.info("Close match in your library - check it before submitting." if top >= 30
                                else "Some related passages found." if top >= 15 else "No strong overlap with your library files.")
                        st.dataframe(pd.DataFrame(res["library"]), hide_index=True)
                    else:
                        st.write("No overlapping passages found.")

    # ---------------- Document ----------------
    with t_doc:
        if comp:
            a, b = st.columns([1, 1.6], gap="large")
            with a:
                st.markdown(card_html("Draft completeness", f"{comp['percent']}%", comp["level"], comp["percent"]), unsafe_allow_html=True)
                st.caption("Measures structure and completeness against the format and checklist, not scientific quality.")
                if audit and audit["limit"]:
                    st.metric("Word count", f"{audit['words']:,} / {audit['limit']:,}",
                              delta="over limit" if audit["words"] > audit["limit"] else "within limit", delta_color="inverse")
                elif audit:
                    st.metric("Word count", f"{audit['words']:,}")
            with b:
                st.markdown('<div class="section-title">To reach 100%</div>', unsafe_allow_html=True)
                if comp["gaps"]:
                    st.markdown("\n".join(f"- {g}" for g in comp["gaps"]))
                else:
                    st.success("Every checklist item is covered - now get your guide's review.")
            st.dataframe(pd.DataFrame(comp["components"]), hide_index=True)
        else:
            st.info("Upload your draft above (before pressing Analyze) to see how complete it is. "
                    "The blueprint below shows what a complete document contains.")
        st.markdown('<div class="section-title">🏁 What the 100% document looks like</div>', unsafe_allow_html=True)
        st.dataframe(pd.DataFrame(res["blueprint"]), hide_index=True)
        if audit:
            st.divider()
            x, y = st.columns(2, gap="large")
            with x:
                if audit.get("sts_sections"):
                    st.markdown('<div class="section-title">ICMR STS format check</div>', unsafe_allow_html=True)
                    st.dataframe(pd.DataFrame(audit["sts_sections"]), hide_index=True)
                    st.caption("Guide counts may be exceeded by up to 100 words (title: 25 max). Font, spacing and "
                               "Vancouver/PubMed-only references can't be checked here.")
                st.markdown('<div class="section-title">Sections found</div>', unsafe_allow_html=True)
                st.dataframe(pd.DataFrame({"Section": list(audit["sections"]),
                                           "In draft": ["✅" if v else "⚠️ missing" for v in audit["sections"].values()]}), hide_index=True)
            with y:
                if audit.get("design_checks"):
                    st.markdown(f'<div class="section-title">Checks for a {escape(audit["design"].lower())}</div>', unsafe_allow_html=True)
                    st.dataframe(pd.DataFrame({"Element": list(audit["design_checks"]),
                                               "In draft": ["✅" if v else "⚠️ missing" for v in audit["design_checks"].values()]}),
                                 hide_index=True)
                if audit.get("identity"):
                    st.error("Remove identifying details (the format says such proposals are rejected): " + ", ".join(audit["identity"]))
                if audit["placeholders"]:
                    st.warning("Placeholder text found: " + ", ".join(audit["placeholders"]))
                st.markdown('<div class="section-title">DOIs cited in your draft</div>', unsafe_allow_html=True)
                if audit["dois"]:
                    st.dataframe(pd.DataFrame(audit["dois"]), hide_index=True)
                else:
                    st.caption("No DOIs found in the draft to verify.")

    # ---------------- AI assessment ----------------
    with t_ai:
        if res.get("design"):
            d = DESIGNS[res["design"]]
            st.markdown(f'<div class="section-title">📘 Design guide: {escape(res["design"])}</div>', unsafe_allow_html=True)
            g1, g2 = st.columns(2, gap="large")
            g1.markdown(f"**Reporting guideline**  \n{d['guideline']}\n\n**Sample size**  \n{d['sample_size']}")
            g2.markdown(f"**Typical analysis**  \n{d['analysis']}\n\n**Ethics route**  \n{d['ethics']}")
            st.caption(f"Fit for a short student project: {d['fit']}. Curated general reference - confirm with your guide and the current call.")
        if ai:
            st.divider()
            if ai.get("gap"):
                st.markdown(f'<div class="card"><div class="lbl">Unexplored gap</div><div style="margin-top:4px">{escape(str(ai["gap"]))}</div></div>',
                            unsafe_allow_html=True)
                st.write("")
            if ai.get("already_done"):
                st.markdown("**Already done in the literature**\n" + "\n".join(f"- {x}" for x in ai["already_done"]))
            m1, m2 = st.columns(2, gap="large")
            with m1:
                st.markdown('<div class="section-title">Suggested methodology</div>', unsafe_allow_html=True)
                meth = ai.get("methodology") or {}
                st.dataframe(pd.DataFrame({"Item": [k.replace("_", " ").title() for k in meth], "Suggestion": [str(v) for v in meth.values()]}),
                             hide_index=True)
            with m2:
                st.markdown('<div class="section-title">Work plan</div>', unsafe_allow_html=True)
                st.dataframe(pd.DataFrame(ai.get("work_plan", [])), hide_index=True)
            e = ai.get("ethics") or {}
            if e:
                st.markdown('<div class="section-title">Ethics screen (advisory - the IEC decides)</div>', unsafe_allow_html=True)
                e1, e2, e3 = st.columns([1, 1, 2])
                e1.markdown(card_html("Human participants", "Yes" if e.get("human_participants") else "No", small=True), unsafe_allow_html=True)
                e2.markdown(card_html("Likely category", e.get("likely_category", "-"), small=True), unsafe_allow_html=True)
                e3.markdown(card_html("Reasoning", e.get("reasons", "-"), small=True), unsafe_allow_html=True)
                st.write("")
            sts = ai.get("sts") or {}
            if sts.get("type_of_study") or sts.get("subject_area"):
                tp, sa = sts.get("type_of_study", ""), sts.get("subject_area", "")
                st.markdown(f"**Suggested STS form entries:** type of study - {tp or '-'}{'' if tp in STS_TYPES else ' (not an official option, check)'}; "
                            f"subject area - {sa or '-'}{'' if sa in STS_SUBJECTS else ' (not an official option, check)'}")
            dr = ai.get("draft_review") or {}
            if any(dr.get(k) for k in ("strengths", "problems", "contradictions")):
                st.markdown('<div class="section-title">AI review of your draft</div>', unsafe_allow_html=True)
                r1, r2, r3 = st.columns(3, gap="large")
                for col, k, icon, title in ((r1, "strengths", "✅", "Strengths"), (r2, "problems", "⚠️", "Problems"),
                                            (r3, "contradictions", "🚨", "Contradictions")):
                    col.markdown(f"**{title}**")
                    for item in dr.get(k, []) or ["-"]:
                        col.markdown(f"{icon} {item}" if item != "-" else "-")
            if ai.get("limits"):
                st.caption(f"Limits noted by the AI: {ai['limits']}")
        else:
            st.info("Add a free Gemini key in the sidebar to unlock the AI assessment: gap analysis, methodology, "
                    "work plan, ethics screen and draft review.")

    # ---------------- Report ----------------
    with t_rep:
        st.markdown('<div class="section-title">Download your report</div>', unsafe_allow_html=True)
        if res.get("pdf"):
            st.download_button("📄 Download official report (PDF)", res["pdf"], file_name="research_assessment_report.pdf",
                               mime="application/pdf", type="primary")
        elif res.get("pdf_error"):
            st.warning(res["pdf_error"])
        st.download_button("📥 Plain-text version (.txt)", res["report"], file_name="research_analysis.txt")
        st.markdown("**The PDF contains:** summary, pentagon chart, novelty analysis, design guide, document completion with the "
                    "100% blueprint, draft checks, closest papers, journal and library comparisons, the AI assessment and a "
                    "limitations section.")
        st.caption(f"Advisory only ({framework}). Novelty depends on the search terms - read the top papers yourself. "
                   "Ethics and exemption decisions rest with your IEC.")


# ---------------------- app ----------------------
def main():
    st.set_page_config(page_title="Apostle9", page_icon="🧠", layout="wide")
    st.markdown(CSS, unsafe_allow_html=True)
    st.markdown(
        '<div class="hero"><div class="t">🧠 Apostle9</div>'
        '<p>Check the novelty of a research idea against real literature, see a five-pillar pentagon assessment, '
        'measure how complete your draft is, and download an official-style PDF report.</p>'
        '<div class="steps"><span class="step">① Describe your idea</span><span class="step">② Add your draft (optional)</span>'
        '<span class="step">③ Analyze</span><span class="step">④ Download the PDF</span></div></div>', unsafe_allow_html=True)

    with st.sidebar:
        st.markdown("### ⚙️ Settings")
        with st.expander("🔑 AI engine (free Gemini key)", expanded=True):
            key = st.text_input("Gemini API key (free at aistudio.google.com)", type="password")
            if not key:
                try:
                    key = st.secrets.get("GEMINI_API_KEY", "")
                except Exception:
                    key = ""
            model = st.text_input("Gemini model", value="gemini-flash-latest").strip()
            key = (key or "").strip()
            st.caption("✅ Key detected" if key else "⚠️ No key detected - paste it, then press Enter")
            if st.button("Test key"):
                if not key:
                    st.error("No key entered yet.")
                else:
                    try:
                        _, used = _gemini_json(key, model, 'Return ONLY this JSON: {"ok": true}')
                        st.success(f"Key works (model: {used})")
                    except Exception as e:
                        st.error(str(e))
            st.caption("No key = rule-based estimates for 4 pillars. Novelty always comes from real search data.")
        with st.expander("📰 Journals to compare"):
            journals = st.multiselect("Compare against journals", DEFAULT_JOURNALS, default=DEFAULT_JOURNALS[:3])
            more_j = st.text_input("Other journals (comma-separated full titles)")
            journals = journals + [j.strip() for j in more_j.split(",") if j.strip()]
        with st.expander("🗂️ Reference library (old ICMR files, papers)"):
            bundled = load_bundled_library()
            st.caption(f"{len(bundled)} file(s) bundled in the app's reference_library folder.")
            lib_up = st.file_uploader("Add files for this session (lists of approved STS titles, past proposals, papers)",
                                      type=["txt", "csv", "pdf", "docx"], accept_multiple_files=True, key="lib")
            session_docs = [{"source": u.name, "text": read_upload(u)} for u in (lib_up or [])]
            library_docs = bundled + [d for d in session_docs if d["text"].strip()]
            st.caption("Compared by shared words. Don't add anyone's confidential proposals, "
                       "and remember a public GitHub repo makes bundled files public.")
        with st.expander("📝 Extra rules and limits"):
            word_limit = st.number_input("Total word limit (optional, 0 = off)", min_value=0, value=0, step=50)
            st.caption("For ICMR STS the official per-section limits are checked automatically.")
            extra_rules = st.text_area("Paste current rules/format to apply (optional)", height=100,
                                       help="E.g. headings and limits from this year's call. These override the built-in notes.")
        with st.expander("📘 Rules & reference"):
            st.caption("ICMR STS rules (official DHR documents, read 22 Oct 2026 - confirm on the portal):")
            for rule, src in STS_RULES:
                st.markdown(f"- {rule} *({src})*")
            st.caption("Concept-note format (DHR 1.1):")
            for name, guide, _ in STS_FORMAT:
                st.markdown(f"- {name}: {guide} words")
            for note in STS_FORMAT_NOTES:
                st.markdown(f"- {note}")
            st.caption("Types of study: " + "; ".join(STS_TYPES))
            st.caption("Subject areas: " + "; ".join(STS_SUBJECTS))
            st.caption("Ethics categories: " + ETHICS_NOTE)
            for k, v in ETHICS_CATEGORIES.items():
                st.markdown(f"- **{k}:** {v}")
            st.caption("Study designs covered: " + ", ".join(DESIGNS))

    left, right = st.columns([3, 2], gap="large")
    with left:
        st.markdown('<div class="section-title">① Describe your idea</div>', unsafe_allow_html=True)
        proposal = st.text_area("Your idea", height=150, label_visibility="collapsed",
                                placeholder="A few keywords is enough (e.g. beta cell, insulin, implant) or a full sentence. "
                                            "You can also leave this blank and upload your draft.")
        custom = st.text_input("Your own keywords (optional, comma-separated - commas keep phrases together)")
    with right:
        st.markdown('<div class="section-title">② Add your draft (optional)</div>', unsafe_allow_html=True)
        up = st.file_uploader("Draft or guideline", type=["txt", "csv", "pdf", "docx"], label_visibility="collapsed")
        extra = read_upload(up)
        if up and extra.strip():
            st.caption(f"✅ Read {len(extra.split()):,} words from {up.name}")
        elif up:
            st.warning("No text found (scanned PDFs are images and can't be read).")
        framework = st.selectbox("Submitting to", ["ICMR STS", "International journal", "General academic"])

    if st.button("🚀 Analyze my idea", type="primary"):
        prop = proposal.strip() or extra[:1500].strip()
        if len(prop) < 3:
            st.warning("Type a few keywords or a sentence about your idea, or upload a draft with text.")
        else:
            try:
                st.session_state["result"] = run_analysis(prop, custom, key, model, extra, word_limit, extra_rules,
                                                         framework, journals, library_docs)
            except Exception as e:
                st.session_state.pop("result", None)
                st.error(f"Analysis failed: {e}")

    if "result" in st.session_state:  # kept in session so the download button doesn't wipe the page
        st.write("")
        render(st.session_state["result"], framework)
    else:
        st.write("")
        e1, e2, e3 = st.columns(3)
        e1.markdown(card_html("Novelty from real papers", "🔎", "Searches Europe PMC with synonyms and checks that each paper is truly relevant", small=True), unsafe_allow_html=True)
        e2.markdown(card_html("Five-pillar pentagon", "⬠", "Novelty, feasibility, pathway, methodology and work plan, each as a percentage", small=True), unsafe_allow_html=True)
        e3.markdown(card_html("Draft completeness + PDF", "📄", "See how close your draft is to 100% and download an official-style report", small=True), unsafe_allow_html=True)


main()
